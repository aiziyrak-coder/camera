"""Oylik davomat tabeli (`/api/hisobot/tabel`).

Dunyo joriy oyga quriladi (tabelda "kelajak kunlari" bo'lishi shart):
  DI-2301: Aliyev (yuzi bor, o'tgan ish kunida 08:05 keldi),
           Botirova (yuzi bor, o'sha kuni 09:15 — kech),
           Nazarov (yuzi YO'Q — hamma kataklari "ma'lumot yo'q")
  DI-2302: Sobirov (yuzi bor, o'sha kuni kelmadi)
"""

from datetime import date, time, timedelta
from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.config import settings
from app.models import AttendanceRecord, Faculty, StudentStaff, User
from app.security import hash_password
from app.services import situation as svc, tabel
from tests.conftest import auth_headers

STEWARD_LOGIN, STEWARD_PASSWORD = "tabel-steward", "Steward123!"


def _person(name, group, faculty, enrolled=True):
    return StudentStaff(full_name=name, type="talaba", group_or_position=group, faculty_id=faculty.id,
                        biometrics_status="tasdiqlangan" if enrolled else "yoq", active=True)


def _month_days(day: date) -> list[date]:
    start = day.replace(day=1)
    end = date(start.year + (start.month == 12), start.month % 12 + 1, 1) - timedelta(days=1)
    return [start + timedelta(days=i) for i in range((end - start).days + 1)]


@pytest.fixture
async def world(db_session, seeded):
    svc.clear_cache()
    today = svc.today()
    days = _month_days(today)
    # Yozuv yoziladigan kun — shu oydagi o'tgan (yoki bugungi) ish kuni.
    work_day = next(d for d in reversed(days) if d <= today and d.isoweekday() in (1, 2, 3, 4, 5, 6))
    di = (await db_session.execute(select(Faculty).where(Faculty.name == "Davolash ishi"))).scalars().one()
    people = [_person("Aliyev Anvar", "2-kurs, DI-2301", di), _person("Botirova Nigora", "2-kurs, DI-2301", di),
              _person("Nazarov Oybek", "2-kurs, DI-2301", di, enrolled=False),
              _person("Sobirov Rustam", "2-kurs, DI-2302", di)]
    db_session.add_all(people)
    await db_session.flush()
    db_session.add_all([
        AttendanceRecord(student_staff_id=people[0].id, date=work_day, status="keldi",
                         check_in=time(8, 5), source="kamera"),
        AttendanceRecord(student_staff_id=people[1].id, date=work_day, status="kech_keldi",
                         check_in=time(9, 15), source="kamera"),
        AttendanceRecord(student_staff_id=people[3].id, date=work_day, status="kelmadi", source="kamera"),
    ])
    await db_session.commit()
    svc.clear_cache()
    yield {"today": today, "days": days, "work_day": work_day, "month": f"{today.year:04d}-{today.month:02d}",
           "faculty": di}
    svc.clear_cache()


@pytest.fixture
async def admin(client: AsyncClient, world):
    return await auth_headers(client, "operator", "operator123")


async def _get(client, headers, **params):
    res = await client.get("/api/hisobot/tabel", params=params, headers=headers)
    assert res.status_code == 200, res.text
    return res.json()


def _row(data, name):
    return next(p for p in data["people"] if p["fullName"] == name)


def _mark(person, day: date) -> dict:
    return next(c for c in person["cells"] if c["day"] == day.day)


# ─────────────────────────────────────────── sof funksiyalar

def test_month_bounds_and_label():
    start, end = tabel.month_bounds("2026-09")
    assert (start, end) == (date(2026, 9, 1), date(2026, 9, 30))
    assert tabel.month_label(start) == "2026-yil sentabr"
    start, end = tabel.month_bounds("2026-12")
    assert end == date(2026, 12, 31)


# ─────────────────────────────────────────── API

async def test_shape_and_days(client, admin, world):
    data = await _get(client, admin, kind="talaba", oy=world["month"])
    assert data["title"] == "Davomat tabeli"
    assert data["month"] == world["month"]
    assert data["monthLabel"].startswith(f"{world['today'].year}-yil ")
    assert len(data["days"]) == len(world["days"])
    assert [d["day"] for d in data["days"]] == [d.day for d in world["days"]]
    assert {m["mark"] for m in data["legend"]} == {"+", "K", "–", "D", "·"}


async def test_weekend_is_off_day(client, admin, world):
    data = await _get(client, admin, kind="talaba", oy=world["month"])
    sunday = next((d for d in world["days"] if d.isoweekday() == 7), None)
    assert sunday is not None  # har oyda kamida bitta yakshanba bor
    meta = next(d for d in data["days"] if d["day"] == sunday.day)
    assert meta["isWorkDay"] is False
    person = _row(data, "Aliyev Anvar")
    assert _mark(person, sunday)["mark"] == "D"
    # "D" kunlar ish kunlari yakuniga kirmaydi
    off_days = sum(1 for d in world["days"] if d.isoweekday() == 7)
    assert person["totals"]["workDays"] == len(world["days"]) - off_days


async def test_present_late_and_absent_marks(client, admin, world):
    data = await _get(client, admin, kind="talaba", oy=world["month"])
    day = world["work_day"]
    ali = _mark(_row(data, "Aliyev Anvar"), day)
    assert ali["mark"] == "+" and "08:05" in ali["title"]
    botir = _mark(_row(data, "Botirova Nigora"), day)
    # 08:00 boshlanish -> 09:15 = 75 daqiqa kech (policy.late_minutes bilan bir xil)
    assert botir["mark"] == "K" and "75 daqiqa kech" in botir["title"]
    assert _row(data, "Botirova Nigora")["totals"]["late"] == 1
    sobir = _mark(_row(data, "Sobirov Rustam"), day)
    assert sobir["mark"] == "–"
    assert _row(data, "Sobirov Rustam")["totals"]["absent"] == 1


async def test_person_without_face_is_still_in_the_sheet(client, admin, world):
    data = await _get(client, admin, kind="talaba", oy=world["month"])
    person = _row(data, "Nazarov Oybek")
    assert person["enrolled"] is False
    work_cells = [c for c in person["cells"] if c["mark"] != "D"]
    assert work_cells and all(c["mark"] == "·" for c in work_cells)
    assert all(c["title"] == "yuzi tizimga kiritilmagan" for c in work_cells)
    assert person["totals"]["absent"] == 0
    assert data["totals"]["notEnrolled"] == 1


async def test_future_days_are_unknown_not_absent(client, admin, world):
    data = await _get(client, admin, kind="talaba", oy=world["month"])
    future = [d for d in world["days"] if d > world["today"] and d.isoweekday() != 7]
    if not future:
        pytest.skip("oyning oxirgi kuni — kelajak kuni yo'q")
    person = _row(data, "Aliyev Anvar")
    cell = _mark(person, future[0])
    assert cell["mark"] == "·" and cell["title"] == "kun hali kelmagan"
    assert next(d for d in data["days"] if d["day"] == future[0].day)["isFuture"] is True


async def test_totals_add_up(client, admin, world):
    data = await _get(client, admin, kind="talaba", oy=world["month"])
    days = len(world["days"])
    for person in data["people"]:
        t = person["totals"]
        off = sum(1 for c in person["cells"] if c["mark"] == "D")
        assert t["present"] + t["late"] + t["absent"] + t["unknown"] == t["workDays"]
        assert t["workDays"] + off == days
    grand = data["totals"]
    assert grand["people"] == len(data["people"]) == 4
    total_off = sum(1 for c in data["people"][0]["cells"] if c["mark"] == "D")
    assert (grand["present"] + grand["late"] + grand["absent"] + grand["unknown"]
            + total_off * grand["people"]) == grand["people"] * days
    assert grand["present"] == 1 and grand["late"] == 1 and grand["absent"] == 1


async def test_scope_filter_and_label(client, admin, world):
    data = await _get(client, admin, kind="talaba", oy=world["month"],
                      faculty=str(world["faculty"].id), course=2, group="DI-2302")
    assert [p["fullName"] for p in data["people"]] == ["Sobirov Rustam"]
    assert data["scope"] == "Davolash ishi, 2-kurs, DI-2302 guruhi"
    assert data["people"][0]["group"] == "DI-2302"


async def test_note_mentions_missing_faces(client, admin, world):
    data = await _get(client, admin, kind="talaba", oy=world["month"])
    # 4 kishidan 1 tasi (25%) — "katta ulush" chegarasidan yuqori
    assert data["note"] and "yuzi tizimga kiritilmagan" in data["note"]


async def test_excel_opens_with_title_block(client, admin, world):
    res = await client.get("/api/hisobot/tabel.xlsx", params={"kind": "talaba", "oy": world["month"]}, headers=admin)
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    assert wb.sheetnames == ["Tabel"]
    ws = wb["Tabel"]
    assert ws["A1"].value == settings.org_name
    assert ws["A2"].value == "Davomat tabeli"
    assert ws["A3"].value == "Barcha talabalar"
    assert ws["A4"].value.endswith(tabel.MONTHS[world["today"].month - 1])
    assert ws["A6"].value == "№" and ws["B6"].value == "F.I.Sh."
    assert ws["D6"].value == 1  # birinchi kun ustuni
    assert ws["A8"].value == 1 and ws["B8"].value == "Aliyev Anvar"
    # Uchta chap ustun qotib turadi (guruh ustuni ham ko'rinib tursin).
    assert ws.freeze_panes == "D8"
    assert ws.page_setup.orientation == "landscape"


async def test_excel_marks_land_on_the_right_day(client, admin, world):
    """Belgilar kun raqami bo'yicha joylanadi, ro'yxatdagi o'rin bo'yicha emas."""
    res = await client.get("/api/hisobot/tabel", params={"kind": "talaba", "oy": world["month"]}, headers=admin)
    data = res.json()
    person = data["people"][0]
    wb = load_workbook(BytesIO((await client.get(
        "/api/hisobot/tabel.xlsx", params={"kind": "talaba", "oy": world["month"]}, headers=admin)).content))
    ws = wb["Tabel"]
    for i, day in enumerate(data["days"]):
        mark = next((c["mark"] for c in person["cells"] if c["day"] == day["day"]), "")
        assert ws.cell(8, 4 + i).value == (mark or None), day["day"]


async def test_permission_is_enforced(client, db_session, seeded, world):
    db_session.add(User(login=STEWARD_LOGIN, password_hash=hash_password(STEWARD_PASSWORD),
                        full_name="Kamera Mas'uli", role="kamera-masuli"))
    await db_session.commit()
    headers = await auth_headers(client, STEWARD_LOGIN, STEWARD_PASSWORD)
    for path in ("/api/hisobot/tabel", "/api/hisobot/tabel.xlsx"):
        res = await client.get(path, params={"kind": "talaba", "oy": world["month"]}, headers=headers)
        assert res.status_code == 403, res.text
    res = await client.get("/api/hisobot/tabel", params={"kind": "talaba", "oy": world["month"]})
    assert res.status_code in (401, 403)


async def test_bad_month_is_rejected(client, admin, world):
    res = await client.get("/api/hisobot/tabel", params={"kind": "talaba", "oy": "2026-13"}, headers=admin)
    assert res.status_code == 422
