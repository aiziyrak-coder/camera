"""Haftalik dars jadvali: namuna fayl va uni semestrga yoyish."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.lesson_weekly import (
    MAX_GENERATED,
    build_weekly_template,
    expand_weekly,
    parse_weekday,
    weekday_dates,
)

pytestmark = pytest.mark.anyio


def test_weekday_is_understood_in_several_spellings():
    assert parse_weekday("Seshanba") == 2
    assert parse_weekday("se") == 2
    assert parse_weekday("SESH") == 2
    assert parse_weekday("2") == 2
    assert parse_weekday("вторник") == 2
    assert parse_weekday("payshanba") == 4
    assert parse_weekday("") is None
    assert parse_weekday("kecha") is None


def test_weekday_dates_cover_the_whole_range():
    # 2026-09-01 — seshanba.
    dates = weekday_dates(2, date(2026, 9, 1), date(2026, 9, 30))
    assert dates[0] == date(2026, 9, 1)
    assert dates[-1] == date(2026, 9, 29)
    assert len(dates) == 5
    # Oraliq boshlanishidan oldin tushadigan kun qo'shilmaydi.
    assert weekday_dates(1, date(2026, 9, 2), date(2026, 9, 5)) == []


def test_expand_creates_one_row_per_date():
    header = ["guruh", "hafta_kuni", "boshlanish", "xona", "fan"]
    rows = [{"guruh": "DI-2301", "hafta_kuni": "Seshanba", "boshlanish": "08:30", "xona": "214", "fan": "Anatomiya"}]
    out = expand_weekly(header, rows, date(2026, 9, 1), date(2026, 9, 30), weekday_column="hafta_kuni")

    assert out.lessons == 5
    assert out.weeks == 5
    assert out.bad_weekday == []
    # Hafta kuni o'rniga sana qo'yiladi — keyin odatdagi import o'qiydi.
    assert "hafta_kuni" not in out.header and "sana" in out.header
    assert out.rows[0]["sana"] == "2026-09-01"
    assert out.rows[0]["guruh"] == "DI-2301"
    assert "hafta_kuni" not in out.rows[0]


def test_unreadable_weekday_is_reported_not_guessed():
    header = ["guruh", "hafta_kuni", "boshlanish"]
    rows = [
        {"guruh": "DI-2301", "hafta_kuni": "Seshanba", "boshlanish": "08:30"},
        {"guruh": "DI-2302", "hafta_kuni": "kecha", "boshlanish": "08:30"},
    ]
    out = expand_weekly(header, rows, date(2026, 9, 1), date(2026, 9, 7), weekday_column="hafta_kuni")
    # Ikkinchi qator tashlab ketilmaydi — fayldagi raqami bilan aytiladi.
    assert out.bad_weekday == [3]
    assert out.lessons == 1


def test_empty_rows_are_skipped():
    header = ["guruh", "hafta_kuni"]
    rows = [{"guruh": "", "hafta_kuni": ""}, {"guruh": "DI-2301", "hafta_kuni": "Juma"}]
    out = expand_weekly(header, rows, date(2026, 9, 1), date(2026, 9, 7), weekday_column="hafta_kuni")
    assert out.lessons == 1
    assert out.bad_weekday == []


def test_huge_range_is_capped_and_says_so():
    header = ["guruh", "hafta_kuni"]
    rows = [{"guruh": f"G-{i}", "hafta_kuni": "Dushanba"} for i in range(200)]
    out = expand_weekly(header, rows, date(2020, 1, 1), date(2026, 1, 1), weekday_column="hafta_kuni")
    assert out.truncated is True
    assert out.lessons == MAX_GENERATED


def test_template_has_the_sheets_and_lists_a_kafedra_needs():
    content = build_weekly_template(
        ["DI-2301", "DI-2302"],
        [("214", "214-xona"), ("215", "215-xona")],
        lesson_minutes=90,
    )
    wb = load_workbook(BytesIO(content))
    assert wb.sheetnames == ["Jadval", "Guruhlar", "Xonalar", "Yo'riqnoma"]

    sheet = wb["Jadval"]
    assert [sheet.cell(1, i).value for i in range(1, 7)] == [
        "Guruh", "Hafta kuni", "Boshlanish", "Xona", "Fan", "O'qituvchi",
    ]
    # Sana ustuni ATAYLAB yo'q: jadval haftalik to'ldiriladi.
    assert "Sana" not in [sheet.cell(1, i).value for i in range(1, 8)]

    assert wb["Guruhlar"].cell(2, 1).value == "DI-2301"
    assert wb["Xonalar"].cell(2, 1).value == "214"
    assert wb["Xonalar"].cell(2, 2).value == "214-xona"
    # Dars davomiyligi aytiladi — kafedra buni o'zi taxmin qilmasin.
    guide = "\n".join(str(cell.value or "") for row in wb["Yo'riqnoma"].iter_rows() for cell in row)
    assert "90 daqiqa" in guide


def test_template_warns_when_no_room_has_a_camera():
    content = build_weekly_template(["DI-2301"], [], lesson_minutes=80)
    wb = load_workbook(BytesIO(content))
    guide = "\n".join(str(cell.value or "") for row in wb["Yo'riqnoma"].iter_rows() for cell in row)
    # Kamerasiz xonadagi darsni tizim tekshira olmaydi — buni yashirmaymiz.
    assert "kamera biriktirilmagan" in guide
