"""Institut kadrlar ro'yxati JSHSHIR bilan yuritiladi, pasport bilan emas.

Ommaviy import qilingan 688 xodimda pasport seriyasi umuman yo'q — ya'ni
ular uchun ro'yxatdan o'tishning yagona yo'li shu raqam. Bu testlar aynan
o'sha yo'lni va uning himoyasini qo'riqlaydi.
"""

import io

import openpyxl
import pytest
from httpx import AsyncClient
from sqlalchemy import select

from app.models import Faculty, StudentStaff
from tests.conftest import auth_headers

PINFL = "30302654150047"


@pytest.fixture
async def a_staff_member(db_session, seeded) -> StudentStaff:
    faculty = (await db_session.execute(select(Faculty))).scalars().first()
    person = StudentStaff(
        full_name="Qayumov G'anisher Olimovich",
        type="xodim",
        pinfl=PINFL,
        faculty_id=faculty.id,
        group_or_position="Xalq tabobati va farmakologiya",
        biometrics_status="yoq",
    )
    db_session.add(person)
    await db_session.commit()
    await db_session.refresh(person)
    return person


@pytest.mark.usefixtures("seeded")
class TestLookupByPinfl:
    async def test_a_staff_member_is_found_by_pinfl(self, client: AsyncClient, a_staff_member):
        resp = await client.post("/api/public/enrollment/lookup", json={"pinfl": PINFL})
        assert resp.status_code == 200
        body = resp.json()
        assert body["fullName"] == "Qayumov G'anisher Olimovich"
        assert body["typeLabel"] == "Xodim"
        assert body["alreadyEnrolled"] is False

    async def test_spaces_and_dashes_are_ignored(self, client: AsyncClient, a_staff_member):
        """Odam raqamni ko'chirib qo'yganda bo'sh joy va chiziqcha qo'shilib
        qolishi juda tez-tez uchraydi. Ularni tozalamasak, raqami to'g'ri
        bo'lgan xodim "topilmadi" javobini olardi va sababini tushunmasdi."""
        resp = await client.post(
            "/api/public/enrollment/lookup", json={"pinfl": " 3030-2654 1500-47 "}
        )
        assert resp.status_code == 200
        assert resp.json()["fullName"] == "Qayumov G'anisher Olimovich"

    async def test_an_unknown_pinfl_is_not_found(self, client: AsyncClient, a_staff_member):
        resp = await client.post("/api/public/enrollment/lookup", json={"pinfl": "99999999999999"})
        assert resp.status_code == 404
        assert "JSHSHIR" in resp.json()["detail"]

    async def test_an_empty_request_is_rejected(self, client: AsyncClient):
        """Na JSHSHIR, na pasport — bu qidiruv emas, xato so'rov."""
        resp = await client.post("/api/public/enrollment/lookup", json={})
        assert resp.status_code == 422

    async def test_the_passport_path_still_works(self, client: AsyncClient, db_session, seeded):
        """Ilgari pasport bilan ro'yxatdan o'tganlar yo'qotilmasligi kerak."""
        person = StudentStaff(
            full_name="Eski Foydalanuvchi", type="talaba", group_or_position="301-guruh",
            passport_series="AD", passport_number="1234567", biometrics_status="yoq",
        )
        db_session.add(person)
        await db_session.commit()

        resp = await client.post(
            "/api/public/enrollment/lookup",
            json={"passportSeries": "AD", "passportNumber": "1234567"},
        )
        assert resp.status_code == 200
        assert resp.json()["fullName"] == "Eski Foydalanuvchi"


@pytest.mark.usefixtures("seeded")
class TestSubmitIsGuardedByTheSameIdentity:
    async def test_a_wrong_pinfl_cannot_upload_a_photo(self, client: AsyncClient, a_staff_member):
        """Yozuv identifikatorini bilgan odam boshqa birovning yozuviga
        rasm yuklay olmasligi kerak. /submit identifikatsiyani /lookup
        bilan aynan bir xil qayta tekshiradi."""
        resp = await client.post(
            f"/api/public/enrollment/{a_staff_member.id}/submit",
            data={"pinfl": "11111111111111"},
            files={"photos": ("f.jpg", b"not-a-real-image", "image/jpeg")},
        )
        assert resp.status_code == 403
        assert "JSHSHIR" in resp.json()["detail"]

    async def test_no_identity_at_all_is_rejected(self, client: AsyncClient, a_staff_member):
        resp = await client.post(
            f"/api/public/enrollment/{a_staff_member.id}/submit",
            files={"photos": ("f.jpg", b"not-a-real-image", "image/jpeg")},
        )
        assert resp.status_code == 422


@pytest.mark.usefixtures("seeded")
class TestBiometricsCoverage:
    async def test_coverage_counts_confirmed_and_missing_per_faculty(
        self, client: AsyncClient, db_session, a_staff_member
    ):
        faculty = (await db_session.execute(select(Faculty))).scalars().first()
        db_session.add_all([
            StudentStaff(
                full_name="Tasdiqlagan Xodim", type="xodim", pinfl="30302654150048",
                faculty_id=faculty.id, group_or_position="Kafedra",
                biometrics_status="tasdiqlangan",
            ),
            StudentStaff(
                full_name="Fakultetsiz Xodim", type="xodim", pinfl="30302654150049",
                group_or_position="Rektorat", biometrics_status="yoq",
            ),
        ])
        await db_session.commit()

        headers = await auth_headers(client, "admin", "admin123")
        resp = await client.get("/api/students-staff/biometrics-coverage?type=xodim", headers=headers)
        assert resp.status_code == 200
        body = resp.json()

        assert body["total"] == 3
        assert body["confirmed"] == 1
        assert body["missing"] == 2

        rows = {r["faculty"]: r for r in body["byFaculty"]}
        assert rows[faculty.name]["total"] == 2
        assert rows[faculty.name]["confirmed"] == 1
        assert rows["Fakultetsiz"]["total"] == 1
        assert rows["Fakultetsiz"]["confirmed"] == 0

    async def test_an_empty_group_reports_no_data_not_zero_percent(
        self, client: AsyncClient, seeded
    ):
        """Hech kim yo'q guruhda "0%" yozish "hech kim tasdiqlamadi"
        degan ma'noni berardi. Bu ikkisi turli xulosaga olib keladi."""
        headers = await auth_headers(client, "admin", "admin123")
        resp = await client.get("/api/students-staff/biometrics-coverage", headers=headers)
        assert resp.status_code == 200
        assert resp.json()["percent"] is None


@pytest.mark.usefixtures("seeded")
class TestFilteringAndExport:
    @pytest.fixture(autouse=True)
    async def _people(self, db_session, seeded):
        faculty = (await db_session.execute(select(Faculty))).scalars().first()
        db_session.add_all([
            StudentStaff(
                full_name="Tasdiqlagan Odam", type="xodim", pinfl="41111111111111",
                faculty_id=faculty.id, group_or_position="Kafedra A",
                biometrics_status="tasdiqlangan",
            ),
            StudentStaff(
                full_name="Tasdiqlamagan Odam", type="xodim", pinfl="42222222222222",
                faculty_id=faculty.id, group_or_position="Kafedra B",
                biometrics_status="yoq",
            ),
        ])
        await db_session.commit()

    async def test_filtering_by_biometrics_status(self, client: AsyncClient):
        headers = await auth_headers(client, "admin", "admin123")
        resp = await client.get("/api/students-staff?biometricsStatus=yoq", headers=headers)
        assert resp.status_code == 200
        names = [r["fullName"] for r in resp.json()["items"]]
        assert "Tasdiqlamagan Odam" in names
        assert "Tasdiqlagan Odam" not in names

    async def test_searching_by_pinfl(self, client: AsyncClient):
        """Kadrlar bo'limi odamni ko'pincha ismi bilan emas, raqami bilan izlaydi."""
        headers = await auth_headers(client, "admin", "admin123")
        resp = await client.get("/api/students-staff?search=42222222222222", headers=headers)
        assert resp.status_code == 200
        assert [r["fullName"] for r in resp.json()["items"]] == ["Tasdiqlamagan Odam"]

    async def _download(self, client, query: str):
        headers = await auth_headers(client, "admin", "admin123")
        resp = await client.get(f"/api/students-staff/export?{query}", headers=headers)
        assert resp.status_code == 200, resp.text
        assert "spreadsheetml" in resp.headers["content-type"]
        assert "attachment" in resp.headers["content-disposition"]
        return openpyxl.load_workbook(io.BytesIO(resp.content))

    async def test_the_people_list_keeps_the_same_filter_as_the_screen(self, client: AsyncClient):
        """Yuklab olingan fayl ekranda ko'rilgan narsaning nusxasi bo'lishi
        kerak — aks holda u hisobot sifatida ishonchsiz."""
        wb = await self._download(client, "kind=people&biometricsStatus=yoq")
        ws = wb["Ro'yxat"]
        names = [ws.cell(row=r, column=2).value for r in range(5, ws.max_row + 1)]
        assert "Tasdiqlamagan Odam" in names
        assert "Tasdiqlagan Odam" not in names
        assert "Tasdiqlanmagan" in ws["A2"].value  # fayl qaysi filtr bilan tuzilganini aytadi

    async def test_columns_are_real_spreadsheet_columns(self, client: AsyncClient):
        """CSV da ustun ajratgichini Excel lokalga qarab TAXMIN qilardi va
        qatorlar bitta katakka aralashib tushardi. .xlsx da har bir qiymat
        o'z katagida."""
        wb = await self._download(client, "kind=people")
        ws = wb["Ro'yxat"]
        header = [ws.cell(row=4, column=c).value for c in range(1, 8)]
        assert header == ["№", "F.I.SH.", "JSHSHIR", "Turi", "Fakultet", "Kafedra / Bo'lim", "Yuz holati"]

    async def test_pinfl_is_stored_as_text_with_leading_zeros(self, client: AsyncClient, db_session):
        """Son sifatida yozilsa Excel 14 xonali raqamni 3,03E+13 qilib
        ko'rsatadi, nol bilan boshlanganini esa shunchaki 0 ga aylantiradi."""
        faculty = (await db_session.execute(select(Faculty))).scalars().first()
        db_session.add(StudentStaff(
            full_name="Nolli Raqam", type="xodim", pinfl="00000000000000",
            faculty_id=faculty.id, group_or_position="Sinov", biometrics_status="yoq",
        ))
        await db_session.commit()

        wb = await self._download(client, "kind=people&search=00000000000000")
        cell = wb["Ro'yxat"].cell(row=5, column=3)
        assert cell.value == "00000000000000"
        assert cell.number_format == "@"

    async def test_the_statistics_file_counts_by_faculty_and_unit(self, client: AsyncClient):
        wb = await self._download(client, "kind=stats")
        assert wb.sheetnames == ["Umumiy statistika", "Kafedra va bo'limlar"]

        summary = wb["Umumiy statistika"]
        values = {summary.cell(row=r, column=1).value: summary.cell(row=r, column=2).value
                  for r in range(4, 9)}
        assert values["Jami ro'yxatda"] == 2
        assert values["Yuzi tasdiqlangan"] == 1
        assert values["Yuzi tasdiqlanmagan"] == 1
        assert values["Qamrov"] == pytest.approx(0.5)

        units = wb["Kafedra va bo'limlar"]
        unit_names = {units.cell(row=r, column=2).value for r in range(5, units.max_row + 1)}
        assert {"Kafedra A", "Kafedra B"} <= unit_names

    async def test_statistics_ignore_the_list_filters(self, client: AsyncClient):
        """"Umumiy statistika" bitta holatga qisqartirilsa o'z nomiga zid
        bo'lardi: yuz holati filtri faqat ro'yxatga tegishli."""
        wb = await self._download(client, "kind=stats&biometricsStatus=yoq")
        summary = wb["Umumiy statistika"]
        assert summary.cell(row=4, column=2).value == 2

    async def test_an_unknown_export_kind_is_rejected(self, client: AsyncClient):
        headers = await auth_headers(client, "admin", "admin123")
        resp = await client.get("/api/students-staff/export?kind=pdf", headers=headers)
        assert resp.status_code == 422

    async def test_export_requires_permission(self, client: AsyncClient):
        resp = await client.get("/api/students-staff/export")
        assert resp.status_code in (401, 403)
