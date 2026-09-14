"""Talabalar va xodimlar alohida: ro'yxat, qamrov va Excel.

Asosiy savollar: "2-kursdan kim ro'yxatdan o'tmadi?" va "xodimlardan kim
o'tdi?". Ikkalasida ham eng qimmat xato — jimgina noto'g'ri javob:
"kutilmoqda" holatidagilarning "o'tmaganlar"dan tushib qolishi, talaba
kursining xodim kafedrasi bilan bitta ustunga aralashishi yoki "1-kurs"
filtri "11-kurs"ni ham ilib ketishi.
"""

import io
from datetime import datetime, timezone

import openpyxl
import pytest
from httpx import AsyncClient
from sqlalchemy import select

from app.models import Faculty, StudentStaff
from app.services.staff_export import split_course
from tests.conftest import auth_headers

CONFIRMED_UTC = datetime(2026, 9, 14, 8, 57, 54, tzinfo=timezone.utc)  # Toshkentda 13:57


@pytest.fixture
async def people(db_session, seeded):
    faculty = (await db_session.execute(select(Faculty).where(Faculty.name == "Davolash ishi"))).scalar_one()
    db_session.add_all([
        StudentStaff(full_name="Aliyev Xodim", type="xodim", faculty_id=faculty.id,
                     group_or_position="Anatomiya kafedrasi", biometrics_status="tasdiqlangan",
                     biometrics_confirmed_at=CONFIRMED_UTC),
        StudentStaff(full_name="Boboyev Xodim", type="xodim", faculty_id=faculty.id,
                     group_or_position="Anatomiya kafedrasi", biometrics_status="yoq"),
        StudentStaff(full_name="Karimova Ikkinchi", type="talaba", faculty_id=faculty.id,
                     group_or_position="2-kurs, DI-1625", biometrics_status="yoq"),
        StudentStaff(full_name="Nazarova Kutayotgan", type="talaba", faculty_id=faculty.id,
                     group_or_position="2-kurs, DI-1625", biometrics_status="kutilmoqda"),
        StudentStaff(full_name="Rahimov Uchinchi", type="talaba", faculty_id=faculty.id,
                     group_or_position="3-kurs, DI-1524", biometrics_status="tasdiqlangan"),
        StudentStaff(full_name="Saidov Fakultetsiz", type="talaba",
                     group_or_position="4-kurs", biometrics_status="yoq"),
        StudentStaff(full_name="Umarov Onbirinchi", type="talaba", faculty_id=faculty.id,
                     group_or_position="11-kurs, X-1", biometrics_status="yoq"),
    ])
    await db_session.commit()


async def _get(client: AsyncClient, path: str, **params):
    headers = await auth_headers(client, "admin", "admin123")
    resp = await client.get(path, params=params, headers=headers)
    assert resp.status_code == 200, resp.text
    return resp


async def _workbook(client: AsyncClient, **params):
    resp = await _get(client, "/api/students-staff/export", **params)
    return openpyxl.load_workbook(io.BytesIO(resp.content))


def _column(ws, header: str) -> list:
    headers = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
    col = headers.index(header) + 1
    return [ws.cell(row=r, column=col).value for r in range(5, ws.max_row + 1)]


class TestSplitCourse:
    @pytest.mark.parametrize(
        ("text", "expected"),
        [
            ("2-kurs, DI-1625", (2, "DI-1625")),
            ("4-kurs", (4, "")),
            ("6-kurs, Kardiologiya-25 (Magistratura)", (6, "Kardiologiya-25 (Magistratura)")),
            ("Anatomiya kafedrasi", (None, "Anatomiya kafedrasi")),
            ("", (None, "")),
        ],
    )
    def test_split(self, text, expected):
        assert split_course(text) == expected


@pytest.mark.usefixtures("people")
class TestList:
    async def test_course_filter_is_exact(self, client: AsyncClient):
        resp = await _get(client, "/api/students-staff", type="talaba", course=1)
        assert resp.json()["items"] == []  # "11-kurs" "1-kurs" deb olinmaydi

        resp = await _get(client, "/api/students-staff", type="talaba", course=2)
        items = resp.json()["items"]
        assert {i["fullName"] for i in items} == {"Karimova Ikkinchi", "Nazarova Kutayotgan"}
        assert {(i["course"], i["group"]) for i in items} == {(2, "DI-1625")}

    async def test_not_registered_includes_pending(self, client: AsyncClient):
        resp = await _get(client, "/api/students-staff", type="talaba", biometricsStatus="tasdiqlanmagan")
        names = {i["fullName"] for i in resp.json()["items"]}
        assert names == {"Karimova Ikkinchi", "Nazarova Kutayotgan", "Saidov Fakultetsiz", "Umarov Onbirinchi"}

    async def test_staff_rows_carry_confirmation_time_not_course(self, client: AsyncClient):
        resp = await _get(client, "/api/students-staff", type="xodim", biometricsStatus="tasdiqlangan")
        (staff,) = resp.json()["items"]
        assert staff["course"] is None
        assert staff["group"] is None
        assert staff["confirmedLabel"] == "14.09.2026 13:57"


@pytest.mark.usefixtures("people")
class TestCoverage:
    async def test_students_by_course(self, client: AsyncClient):
        body = (await _get(client, "/api/students-staff/biometrics-coverage", type="talaba")).json()
        assert body["total"] == 5
        rows = {r["course"]: r for r in body["byCourse"]}
        assert list(rows) == ["2-kurs", "3-kurs", "4-kurs", "11-kurs"]
        assert rows["2-kurs"]["total"] == 2
        assert rows["2-kurs"]["courseNumber"] == 2
        assert rows["3-kurs"]["confirmed"] == 1

    async def test_staff_have_no_course_rows(self, client: AsyncClient):
        body = (await _get(client, "/api/students-staff/biometrics-coverage", type="xodim")).json()
        assert body["total"] == 2
        assert body["byCourse"] == []


@pytest.mark.usefixtures("people")
class TestPeopleExport:
    async def test_students_have_their_own_columns(self, client: AsyncClient):
        wb = await _workbook(client, kind="people", type="talaba", course=2, biometricsStatus="tasdiqlanmagan")
        ws = wb["Talabalar"]
        headers = [ws.cell(row=4, column=c).value for c in range(1, 9)]
        assert headers == ["№", "F.I.SH.", "JSHSHIR", "Fakultet", "Kurs", "Guruh", "Yuz holati", "Tasdiqlagan vaqti"]
        assert _column(ws, "F.I.SH.") == ["Karimova Ikkinchi", "Nazarova Kutayotgan"]
        assert set(_column(ws, "Kurs")) == {"2-kurs"}
        assert set(_column(ws, "Guruh")) == {"DI-1625"}
        assert "ro'yxatdan o'tmaganlar" in ws["A1"].value
        assert "2-kurs" in ws["A2"].value

    async def test_staff_have_their_own_columns_and_time(self, client: AsyncClient):
        wb = await _workbook(client, kind="people", type="xodim")
        ws = wb["Xodimlar"]
        assert "Kurs" not in [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
        assert _column(ws, "F.I.SH.") == ["Aliyev Xodim", "Boboyev Xodim"]
        assert _column(ws, "Tasdiqlagan vaqti") == ["14.09.2026 13:57", "—"]
        assert _column(ws, "Kafedra / Bo'lim") == ["Anatomiya kafedrasi", "Anatomiya kafedrasi"]

    async def test_students_are_sorted_by_faculty_course_group_and_nobody_else_leaks_in(self, client: AsyncClient):
        wb = await _workbook(client, kind="people", type="talaba")
        ws = wb["Talabalar"]
        assert _column(ws, "F.I.SH.") == [
            "Karimova Ikkinchi", "Nazarova Kutayotgan",  # Davolash ishi, 2-kurs
            "Rahimov Uchinchi",  # 3-kurs
            "Umarov Onbirinchi",  # 11-kurs
            "Saidov Fakultetsiz",  # fakultetsizlar oxirida
        ]

    async def test_registered_only(self, client: AsyncClient):
        wb = await _workbook(client, kind="people", type="talaba", biometricsStatus="tasdiqlangan")
        assert _column(wb["Talabalar"], "F.I.SH.") == ["Rahimov Uchinchi"]
        assert "ro'yxatdan o'tganlar" in wb["Talabalar"]["A1"].value

    async def test_unknown_type_is_rejected(self, client: AsyncClient):
        headers = await auth_headers(client, "admin", "admin123")
        resp = await client.get("/api/students-staff/export", params={"type": "mehmon"}, headers=headers)
        assert resp.status_code == 422


@pytest.mark.usefixtures("people")
class TestStatsExport:
    async def test_student_statistics_by_course_and_group(self, client: AsyncClient):
        wb = await _workbook(client, kind="stats", type="talaba")
        assert wb.sheetnames == ["Umumiy statistika", "Kurslar", "Guruhlar"]
        courses = wb["Kurslar"]
        assert _column(courses, "Kurs") == ["2-kurs", "3-kurs", "4-kurs", "11-kurs"]
        assert _column(courses, "Jami") == [2, 1, 1, 1]
        groups = wb["Guruhlar"]
        assert ("Davolash ishi", "2-kurs", "DI-1625") in list(
            zip(_column(groups, "Fakultet"), _column(groups, "Kurs"), _column(groups, "Guruh"))
        )

    async def test_staff_statistics_by_department(self, client: AsyncClient):
        wb = await _workbook(client, kind="stats", type="xodim")
        assert wb.sheetnames == ["Umumiy statistika", "Kafedra va bo'limlar"]
        summary = wb["Umumiy statistika"]
        assert summary.cell(row=4, column=2).value == 2
