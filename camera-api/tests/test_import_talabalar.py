"""Talabalar importi: ma'lumotni tozalash qoidalari va bazaga yozish.

Import serverda qo'lda ishga tushiriladi va 6 mingga yaqin odamga
tegadi. Eng qimmat xatolar — jimgina sodir bo'ladiganlari: noto'g'ri
kurs, xodim yozuvini talabaga aylantirish, qayta ishga tushirilganda
yuz tasdig'ini o'chirish, yoki bitta pasportni ikki yozuvga berish
(ro'yxatdan o'tish sahifasi shunda xato beradi).
"""

import openpyxl
import pytest
from sqlalchemy import func, select

from app.models import Faculty, StudentGroup, StudentStaff
from scripts.import_talabalar import (
    FACULTY_DAVOLASH,
    FACULTY_PEDIATRIYA,
    FACULTY_XALQARO,
    build_records,
    clean_group,
    clean_pinfl,
    course_from_filename,
    export_excel,
    faculty_for,
    format_name,
    parse_passport,
    run,
)
from tests.conftest import TestSessionLocal


def _row(pinfl="60000000000011", name="NAZAROVA DILNOZA OYBEK QIZI", passport="AA1000001",
         group="DI-1621", faculty="", file_course=1, source="t.xlsx / 1"):
    return {"pinfl": pinfl, "name": name, "passport": passport, "group": group,
            "faculty": faculty, "file_course": file_course, "source": source}


class TestNames:
    @pytest.mark.parametrize(
        ("raw", "expected"),
        [
            ("TO‘XTAYEVA MALIKA SHERZOD QIZI", "To'xtayeva Malika Sherzod qizi"),
            ("SAIDOV BOBUR ANVAR O‘G‘LI", "Saidov Bobur Anvar o'g'li"),
            ("KARIMOVA RO`ZA NODIR QIZI", "Karimova Ro'za Nodir qizi"),
            ("G‘ANIYEV AHMADJON", "G'aniyev Ahmadjon"),
            ("XXX SARA BINT AHMAD XXX", "Sara Bint Ahmad"),
            ("KHAN  AMIRA RASHID XXX", "Khan Amira Rashid"),
            ("ABDUL-AZIZ KARIMOV", "Abdul-Aziz Karimov"),
            ("JUMABAEVA AYGUL ALÍM QÍZÍ", "Jumabaeva Aygul Alím qízí"),
        ],
    )
    def test_formatting(self, raw, expected):
        assert format_name(raw) == expected

    def test_a_name_of_only_placeholders_is_empty(self):
        assert format_name("XXX XXX") == ""


class TestFields:
    def test_pinfl_stored_as_number_in_excel_keeps_all_digits(self):
        assert clean_pinfl(60000000000011) == "60000000000011"
        assert clean_pinfl(60000000000011.0) == "60000000000011"
        assert clean_pinfl(" 6000-0000 000011 ") == "60000000000011"

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [
            ("AA1000006", ("AA", "1000006")),
            ("aa 100 0006", ("AA", "1000006")),
            ("FA-AB1234567", ("FAAB", "1234567")),
            ("Y1000003", (None, None)),  # bir harfli xorijiy pasport — sahifa qabul qilmaydi
            ("", (None, None)),
        ],
    )
    def test_passport(self, raw, expected):
        assert parse_passport(raw) == expected

    def test_group_spacing(self):
        assert clean_group("ЛД - 25101") == "ЛД-25101"
        assert clean_group("P-2621 A ") == "P-2621 A"

    @pytest.mark.parametrize(
        ("group", "expected"),
        [
            ("DI-1621", FACULTY_DAVOLASH),
            ("ЛД-25101", FACULTY_DAVOLASH),
            ("MD-101/25", FACULTY_XALQARO),
            ("PI-2522", FACULTY_PEDIATRIYA),
            ("P-5024A", FACULTY_PEDIATRIYA),
            ("S-8124", FACULTY_PEDIATRIYA),
            ("Kardiologiya-25", None),
            ("Sog'liqni saqlash-24", None),
            ("BM-1125", None),
            ("", None),
        ],
    )
    def test_faculty_from_group_prefix(self, group, expected):
        assert faculty_for(group) == expected

    def test_an_explicit_faculty_column_wins(self):
        assert faculty_for("XYZ-1", "Pediatriya fakulteti") == FACULTY_PEDIATRIYA

    @pytest.mark.parametrize(("name", "course"), [("1-kurs.xlsx", 1), ("3- kurs.xlsx", 3), ("4 kurslar.xlsx", 4)])
    def test_course_from_filename(self, name, course):
        assert course_from_filename(name) == course


class TestBuildRecords:
    def test_everyone_moves_up_one_course(self):
        """"1-kurs" fayli o'tgan o'quv yiliniki — bu talabalar hozir 2-kursda."""
        (record,), _ = build_records([_row(file_course=1)])
        assert record.course == 2
        assert record.group_or_position == "2-kurs, DI-1621"

    def test_the_last_file_becomes_sixth_year(self):
        (record,), _ = build_records([_row(file_course=5)])
        assert record.course == 6

    def test_without_a_group_only_the_course_is_shown(self):
        (record,), _ = build_records([_row(group="", file_course=3)])
        assert record.group_or_position == "4-kurs"
        assert record.faculty is None

    def test_bad_and_duplicate_rows_are_reported_not_imported(self):
        records, problems = build_records([
            _row(pinfl="123", source="a"),
            _row(name="XXX", source="b"),
            _row(source="c"),
            _row(source="d"),  # c bilan bir xil JSHSHIR
        ])
        assert len(records) == 1
        assert len(problems) == 3


@pytest.mark.usefixtures("seeded")
class TestImportRun:
    ROWS = [
        _row(pinfl="60000000000011", name="NAZAROVA DILNOZA OYBEK QIZI", passport="AA1000001", group="DI-1621", file_course=5),
        _row(pinfl="50000000000022", name="RAHIMOV JASUR BAHODIR O‘G‘LI", passport="AA1000002", group="DI-1621", file_course=5),
        _row(pinfl="60000000000033", name="XXX SARA BINT AHMAD XXX", passport="Y1000003", group="MD-101/25", file_course=1),
        _row(pinfl="50000000000044", name="O‘SAROV ULUG‘BEK OLIM O‘G‘LI", passport="AA1000004", group="", file_course=3),
        _row(pinfl="60000000000055", name="ERGASHEVA NILUFAR TOHIR QIZI", passport="AA1000005", group="Kardiologiya-25", file_course=1),
    ]

    async def _students(self, db_session):
        rows = (await db_session.execute(select(StudentStaff).where(StudentStaff.type == "talaba"))).scalars().all()
        return {r.pinfl: r for r in rows}

    async def test_every_student_reaches_the_database(self, db_session):
        assert await run(self.ROWS, session_factory=TestSessionLocal) == 0
        students = await self._students(db_session)
        assert len(students) == 5

        nazarova = students["60000000000011"]
        assert nazarova.full_name == "Nazarova Dilnoza Oybek qizi"
        assert nazarova.group_or_position == "6-kurs, DI-1621"
        assert (nazarova.passport_series, nazarova.passport_number) == ("AA", "1000001")
        assert nazarova.biometrics_status == "yoq"

        faculty = await db_session.get(Faculty, nazarova.faculty_id)
        assert faculty.name == FACULTY_DAVOLASH  # seed'dagi "Davolash ishi" qayta ishlatildi va nomi aniqlashtirildi

        assert students["50000000000044"].faculty_id is None
        assert students["60000000000055"].faculty_id is None
        assert students["60000000000033"].passport_number is None

    async def test_running_twice_changes_nothing(self, db_session):
        await run(self.ROWS, session_factory=TestSessionLocal)
        first = await db_session.scalar(select(func.count()).select_from(StudentStaff))
        await run(self.ROWS, session_factory=TestSessionLocal)
        second = await db_session.scalar(select(func.count()).select_from(StudentStaff))
        assert first == second

    async def test_a_confirmed_face_survives_a_second_run(self, db_session):
        await run(self.ROWS, session_factory=TestSessionLocal)
        record = (await db_session.execute(
            select(StudentStaff).where(StudentStaff.pinfl == "60000000000011"))).scalar_one()
        record.biometrics_status = "tasdiqlangan"
        record.biometric_embedding = "[0.1]"
        await db_session.commit()

        await run(self.ROWS, session_factory=TestSessionLocal)
        await db_session.refresh(record)
        assert record.biometrics_status == "tasdiqlangan"
        assert record.biometric_embedding == "[0.1]"

    async def test_a_staff_member_with_the_same_pinfl_is_left_alone(self, db_session):
        """Masalan, bir vaqtda institutda ishlaydigan ordinator."""
        db_session.add(StudentStaff(full_name="Nazarova Dilnoza", type="xodim", pinfl="60000000000011",
                                    group_or_position="Kardiologiya", biometrics_status="tasdiqlangan"))
        await db_session.commit()

        await run(self.ROWS, session_factory=TestSessionLocal)
        record = (await db_session.execute(
            select(StudentStaff).where(StudentStaff.pinfl == "60000000000011"))).scalar_one()
        assert record.type == "xodim"
        assert record.group_or_position == "Kardiologiya"

    async def test_a_passport_already_used_elsewhere_is_not_duplicated(self, db_session):
        """Ikkita yozuvda bir pasport bo'lsa, sahifadagi pasport qidiruvi
        bittadan ortiq natija topib, xato beradi."""
        db_session.add(StudentStaff(full_name="Boshqa Odam", type="xodim", pinfl="30000000000001",
                                    passport_series="AA", passport_number="1000001",
                                    group_or_position="Bo'lim", biometrics_status="yoq"))
        await db_session.commit()

        await run(self.ROWS, session_factory=TestSessionLocal)
        owners = await db_session.scalar(
            select(func.count()).select_from(StudentStaff)
            .where(StudentStaff.passport_series == "AA").where(StudentStaff.passport_number == "1000001"))
        assert owners == 1

    async def test_groups_and_faculty_counts_come_from_the_real_list(self, db_session):
        await run(self.ROWS, session_factory=TestSessionLocal)
        groups = {g.name: g for g in (await db_session.execute(select(StudentGroup))).scalars().all()}
        assert groups["DI-1621"].course == 6
        assert groups["DI-1621"].student_count == 2
        assert "Kardiologiya-25" not in groups  # fakulteti noma'lum — guruh jadvaliga yozilmaydi

        faculty = await db_session.get(Faculty, groups["DI-1621"].faculty_id)
        await db_session.refresh(faculty)
        assert faculty.student_count == 2
        assert faculty.course_count == 6

    async def test_dry_run_writes_nothing(self, db_session):
        await run(self.ROWS, dry_run=True, session_factory=TestSessionLocal)
        assert await self._students(db_session) == {}


def test_export_reads_differently_ordered_columns(tmp_path):
    """Har bir kurs fayli ustunlarni boshqacha tartibda saqlaydi, 3-kurs
    faylida esa ismsiz qo'shimcha varaq bor."""
    first = openpyxl.Workbook()
    ws = first.active
    ws.append(["To‘liq ismi", "Pasport raqami", "Talaba ID", "JSHSHIR-kod", "Guruh", "Rasm"])
    ws.append(["NAZAROVA DILNOZA", "AA1000001", "344250000001", 60000000000011, "TPI-925", "#VALUE!"])
    first.save(tmp_path / "1-kurs.xlsx")

    third = openpyxl.Workbook()
    empty = third.active
    empty.title = "xalqaro"
    empty.append(["Talaba ID", "To‘liq ismi", "Pasport raqami", "JSHSHIR-kod", "Guruh", "Foto"])
    empty.append([None, None, None, None, "MD-231", None])
    ws = third.create_sheet("Talabalar")
    ws.append(["To‘liq ismi", "Talaba ID", "JSHSHIR-kod", "Pasport raqami", "Foto"])
    ws.append(["O‘SAROV ULUG‘BEK", "344230000002", "50000000000044", "AA1000004", "#VALUE!"])
    third.save(tmp_path / "3- kurs.xlsx")

    payload = export_excel(tmp_path, tmp_path / "out.json")
    rows = {r["pinfl"]: r for r in payload["rows"]}
    assert set(rows) == {"60000000000011", "50000000000044"}
    assert rows["60000000000011"]["group"] == "TPI-925"
    assert rows["60000000000011"]["file_course"] == 1
    assert rows["50000000000044"]["file_course"] == 3
    assert rows["50000000000044"]["group"] == ""
