"""Hisobotlar (`/api/hisobot`) — tests/situation_world.py dunyosida.

Bugungi talaba davomati (faollar): keldi 6 (Botirova kechikib), kelmadi 2.
Davolash ishi: keldi 5, kelmadi 1. DI-2301 kecha+bugun: keldi 3, kelmadi 3.
Darslar (bugun): L1 — keldi 1, kech 1, kelmadi 2 (L2 hali baholanmagan).
O'qituvchilar (kecha+bugun): Yusupova L1 o'z vaqtida, LY kelmagan; Rahimov L2, L6 kelmagan.
"""

from datetime import datetime, time, timedelta, timezone

import pytest
from httpx import AsyncClient

from app.models import Event
from app.services import hisobot
from tests.conftest import auth_headers
from tests.situation_world import _record, _situation_settings, world  # noqa: F401 — pytest fikstura


@pytest.fixture
async def admin(client: AsyncClient, world):
    return await auth_headers(client, "operator", "operator123")


async def _get(client, headers, **params):
    res = await client.get("/api/hisobot/report", params=params, headers=headers)
    assert res.status_code == 200, res.text
    return res.json()


def _ind(data, key):
    return next(c for c in data["criteria"] if c["key"] == key)["indicator"]


# ─────────────────────────────────────────── sof funksiyalar

def _m(name, raw, faculty=None):
    import uuid
    return hisobot.Member(uuid.uuid4(), name, raw, faculty, None, True)


def test_filter_members_student_cascade():
    import uuid
    fac = uuid.uuid4()
    ms = [_m("Ali", "2-kurs, DI-2301", fac), _m("Vali", "2-kurs, DI-2302", fac), _m("Soli", "3-kurs, DI-2101", fac),
          _m("Guli", "2-kurs, DI-2301", None)]
    f = hisobot.Filters(faculty=str(fac), course=2)
    assert [m.name for m in hisobot.filter_members(ms, "talaba", f)] == ["Ali", "Vali"]
    f = hisobot.Filters(faculty=str(fac), group=" di-2301 ")
    assert [m.name for m in hisobot.filter_members(ms, "talaba", f)] == ["Ali"]
    f = hisobot.Filters(faculty="none")
    assert [m.name for m in hisobot.filter_members(ms, "talaba", f)] == ["Guli"]
    f = hisobot.Filters(q="VAL")
    assert [m.name for m in hisobot.filter_members(ms, "talaba", f)] == ["Vali"]


def test_criteria_are_separate_per_kind():
    student = {c.key for c in hisobot.criteria_for("talaba")}
    staff = {c.key for c in hisobot.criteria_for("xodim")}
    assert {"davomat", "kechikish", "dars_qatnashish", "uxlash"} == student
    assert {"davomat", "kechikish", "erta_ketish", "dars_otkazish", "forma", "tashqari_kirish"} == staff


# ─────────────────────────────────────────── API

async def test_student_attendance_all_and_faculty(client, admin, world):
    data = await _get(client, admin, kind="talaba")
    assert data["criterion"] == "davomat"
    assert data["population"]["total"] == 12  # nofaol hisobga kirmaydi
    assert _ind(data, "davomat") == "75%"
    assert _ind(data, "kechikish") == "1"
    assert data["report"]["breakdown"]["title"] == "Fakultetlar bo'yicha"
    # xodimlar mezoni talabada yo'q
    assert all(c["key"] not in ("erta_ketish", "forma") for c in data["criteria"])

    data = await _get(client, admin, kind="talaba", faculty=str(world.di.id))
    rows = {r["name"]: r["value"] for r in data["report"]["breakdown"]["rows"]}
    assert data["report"]["breakdown"]["title"] == "Kurslar bo'yicha"
    assert rows["2-kurs"] == 75.0 and rows["3-kurs"] == 100.0


async def test_student_group_people_worst_first(client, admin, world):
    data = await _get(client, admin, kind="talaba", faculty=str(world.di.id), course=2, group="DI-2301",
                      **{"from": world.yesterday.isoformat(), "to": world.today.isoformat()})
    body = data["report"]
    assert _ind(data, "davomat") == "50%"
    names = [p["full_name"] for p in body["people"]]
    assert names[:3] == ["Choriyev Sardor", "Botirova Nigora", "Aliyev Anvar"]
    assert body["people"][1]["values"]["rate"] == 50.0
    assert len(body["trend"]["points"]) == 2


async def test_student_lessons_and_sleep(client, admin, world, db_session):
    db_session.add(Event(camera_name="205-xona", building="Bosh bino", module_code=20, module_name="Uxlash",
                         group="E", confidence=80, severity="past", status="yangi", person_name="Aliyev Anvar",
                         occurred_at=datetime.now(timezone.utc)))
    await db_session.commit()
    data = await _get(client, admin, kind="talaba", criterion="dars_qatnashish")
    assert _ind(data, "dars_qatnashish") == "50%"
    assert _ind(data, "uxlash") == "1"
    data = await _get(client, admin, kind="talaba", criterion="uxlash")
    assert [p["full_name"] for p in data["report"]["people"]] == ["Aliyev Anvar"]
    # xodimlarda bu signal ko'rinmaydi
    staff = await _get(client, admin, kind="xodim")
    assert all(c["key"] != "uxlash" for c in staff["criteria"])


async def test_staff_punctuality_lateness_and_units(client, admin, world):
    params = {"from": world.yesterday.isoformat(), "to": world.today.isoformat()}
    data = await _get(client, admin, kind="xodim", criterion="dars_otkazish", **params)
    assert _ind(data, "dars_otkazish") == "25%"
    assert [p["full_name"] for p in data["report"]["people"]] == ["Rahimov Bobur", "Yusupova Dilnoza Anvarovna"]

    data = await _get(client, admin, kind="xodim", criterion="kechikish")
    assert data["report"]["people"][0]["full_name"] == "Karimov Aziz Olimovich"
    assert data["report"]["people"][0]["values"]["late_min"] == 80  # bitta kun — daqiqada

    options = (await client.get("/api/hisobot/filters", params={"kind": "xodim"}, headers=admin)).json()
    anatomy = next(u for u in options["units"] if u["name"] == "Anatomiya kafedrasi")
    assert anatomy["count"] == 2
    data = await _get(client, admin, kind="xodim", unit=anatomy["id"])
    assert data["population"]["total"] == 2
    data = await _get(client, admin, kind="xodim", unit_kind="lavozim")
    assert data["population"]["total"] == 1


async def test_staff_early_leave_and_coat(client, admin, world, db_session):
    today = world.today
    monday = today - timedelta(days=today.weekday() + 7)  # o'tgan haftaning dushanbasi
    db_session.add(_record(world.people.rahimov, monday, "keldi", "07:50", "15:00"))
    db_session.add(Event(camera_name="Kirish-1", building="Bosh bino", module_code=10, module_name="Oq xalat",
                         group="C", confidence=50, severity="past", status="tasdiqlangan",
                         occurred_at=datetime.combine(monday, time(10, 0), tzinfo=timezone.utc)))
    await db_session.commit()
    params = {"from": monday.isoformat(), "to": today.isoformat()}
    data = await _get(client, admin, kind="xodim", criterion="erta_ketish", **params)
    assert _ind(data, "erta_ketish") == "1"
    assert data["report"]["people"][0]["full_name"] == "Rahimov Bobur"

    data = await _get(client, admin, kind="xodim", criterion="forma", **params)
    body = data["report"]
    assert body["people"] == [] and body["columns"] == []
    assert body["breakdown"]["rows"] == [{"id": "Bosh bino", "name": "Bosh bino", "value": 1, "detail": None,
                                          "headcount": None}]


async def test_student_filters_and_export(client, admin, world):
    options = (await client.get("/api/hisobot/filters", params={"kind": "talaba"}, headers=admin)).json()
    assert {f["name"] for f in options["faculties"]} >= {"Davolash ishi", "Pediatriya", "Fakultetsiz"}
    assert any(g["name"] == "DI-2301" and g["course"] == 2 for g in options["groups"])

    res = await client.get("/api/hisobot/export.xlsx", params={"kind": "talaba", "criterion": "davomat"},
                           headers=admin)
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert res.content[:2] == b"PK"


# ─────────────────────────────────────────── tushunarlilik: gap, izoh, bo'sh holat

async def test_summary_sentence_matches_the_tiles(client, admin, world):
    """Tepadagi gap plitkalardagi ayni sonlardan tuziladi."""
    data = await _get(client, admin, kind="talaba")
    body = data["report"]
    assert data["scope"] == "Barcha talabalar"
    first = body["summary"][0]
    # bugun: 6 keldi (1 tasi kech), 2 kelmadi, 4 kishida yozuv yo'q
    assert "12 talabadan" in first and "bugun" in first
    assert "5 tasi o'z vaqtida keldi" in first
    assert "1 tasi kech keldi" in first
    assert "2 tasi kelmadi" in first
    assert "4 tasi haqida yozuv yo'q" in first
    assert "Kelganlar ulushi — 75%." in body["summary"]
    # yuzi kiritilmaganlar haqida ogohlantirish (12 dan 3 tasi)
    assert any("3 tasining yuzi tizimga kiritilmagan" in line for line in body["summary"])
    assert data["population"] == {"total": 12, "enrolled": 9, "not_enrolled": 3}


async def test_criteria_explain_themselves(client, admin, world):
    data = await _get(client, admin, kind="xodim")
    by_key = {c["key"]: c for c in data["criteria"]}
    assert all(c["description"] for c in data["criteria"])
    # vaqtlar kodda emas, ish vaqti sozlamasidan (08:00 + 10 daqiqa)
    assert "08:10" in by_key["kechikish"]["description"]
    assert "17:00" in by_key["erta_ketish"]["description"]
    assert all(c["available"] is (c["unavailable"] is None) for c in data["criteria"])


async def test_blocked_criterion_says_why_instead_of_zero(client, admin, world):
    """Dars jadvali yo'q kunda — "0%" emas, sabab."""
    old = (world.today - timedelta(days=60)).isoformat()
    data = await _get(client, admin, kind="talaba", criterion="dars_qatnashish", **{"from": old, "to": old})
    assert _ind(data, "dars_qatnashish") == "—"
    body = data["report"]
    assert body["blocked"] is True
    assert "dars jadvali" in body["note"]
    assert body["empty"]["description"]
    assert body["summary"][1] == body["note"]


async def test_day_table_has_reader_friendly_columns(client, admin, world):
    """Bitta kun tanlanganda jadval: holat, vaqtlar, kechikish, izoh."""
    data = await _get(client, admin, kind="talaba")
    body = data["report"]
    assert [c["key"] for c in body["columns"]] == ["holat", "check_in", "check_out", "late_min", "note"]
    assert [c["label"] for c in body["columns"]][:2] == ["Holati", "Kelgan vaqti"]
    assert body["people_total"] == 12  # yozuvi yo'qlar ham ko'rinadi
    rows = {p["full_name"]: p["values"] for p in body["people"]}
    assert rows["Botirova Nigora"]["holat"] == "Kech keldi"
    assert rows["Botirova Nigora"]["check_in"] == "09:15"
    assert rows["Botirova Nigora"]["late_min"] == 75
    assert rows["Choriyev Sardor"]["holat"] == "Kelmadi"
    # yuzi kiritilmagan odam uchun sabab yoziladi, bo'sh katak emas
    assert "Yuzi tizimga kiritilmagan" in rows["Ergasheva Laylo"]["note"]
    assert body["people"][0]["values"]["holat"] == "Kelmadi"  # eng muammolisi birinchi
    assert "Tartib:" in body["people_hint"]


async def test_export_columns_match_the_table(client, admin, world):
    from io import BytesIO

    from openpyxl import load_workbook

    data = await _get(client, admin, kind="talaba")
    res = await client.get("/api/hisobot/export.xlsx", params={"kind": "talaba"}, headers=admin)
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    rows = list(wb["Odamlar"].iter_rows(values_only=True))
    header = [c for c in rows[3] if c is not None]
    expected = ["№", "F.I.Sh.", "Guruh yoki bo'linma"] + [
        f"{c['label']}, {c['unit']}" if c["unit"] else c["label"] for c in data["report"]["columns"]]
    assert header == expected
    assert wb["Hisobot"]["B2"].value == data["scope"]


async def test_requires_auth(client, world):
    res = await client.get("/api/hisobot/report", params={"kind": "talaba"})
    assert res.status_code == 401


