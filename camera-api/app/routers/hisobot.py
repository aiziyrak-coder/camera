"""Hisobotlar sahifasi (`/api/hisobot/*`): xodimlar va talabalar alohida.

  GET /api/hisobot/filters?kind=talaba|xodim      — filtr variantlari
  GET /api/hisobot/report?kind=&from=&to=&criterion=&faculty=&course=&group=&unit_kind=&unit=&q=
  GET /api/hisobot/export.xlsx (report bilan bir xil parametrlar) — Excel
  GET /api/hisobot/tabel?kind=&oy=YYYY-MM&<filtrlar>  — oylik davomat tabeli
  GET /api/hisobot/tabel.xlsx (tabel bilan bir xil parametrlar) — chop etish uchun

Hisob-kitob app/services/hisobot.py da.
"""

from io import BytesIO
from typing import Annotated, Literal
from urllib.parse import quote

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.database import get_db
from app.dependencies import CurrentUser, require_permission
from app.services import hisobot, situation as svc, tabel as tabel_svc

router = APIRouter(prefix="/api/hisobot", tags=["hisobot"])

ReadDep = Annotated[CurrentUser, Depends(require_permission("viewReports"))]
DbDep = Annotated[AsyncSession, Depends(get_db)]
KindQuery = Annotated[Literal["talaba", "xodim"], Query()]


def _filters(faculty: str | None, course: int | None, group: str | None, unit_kind: str | None, unit: str | None,
             q: str | None) -> hisobot.Filters:
    clean = lambda v: (v or "").strip() or None  # noqa: E731
    return hisobot.Filters(clean(faculty), course, clean(group), clean(unit_kind), clean(unit), clean(q))


@router.get("/filters")
async def filters(db: DbDep, _: ReadDep, kind: KindQuery = "xodim") -> dict:
    return await hisobot.filter_options(db, kind)


@router.get("/report")
async def report(
    db: DbDep, _: ReadDep, kind: KindQuery = "xodim",
    date_from: Annotated[str | None, Query(alias="from")] = None,
    date_to: Annotated[str | None, Query(alias="to")] = None,
    criterion: str | None = None, faculty: str | None = None, course: Annotated[int | None, Query(ge=1, le=7)] = None,
    group: str | None = None, unit_kind: str | None = None, unit: str | None = None, q: str | None = None,
) -> dict:
    start, end = svc.resolve_range(date_from, date_to, default_days=1)
    return await hisobot.report(db, kind, start, end, _filters(faculty, course, group, unit_kind, unit, q), criterion)


@router.get("/export.xlsx")
async def export(
    db: DbDep, _: ReadDep, kind: KindQuery = "xodim",
    date_from: Annotated[str | None, Query(alias="from")] = None,
    date_to: Annotated[str | None, Query(alias="to")] = None,
    criterion: str | None = None, faculty: str | None = None, course: Annotated[int | None, Query(ge=1, le=7)] = None,
    group: str | None = None, unit_kind: str | None = None, unit: str | None = None, q: str | None = None,
) -> Response:
    start, end = svc.resolve_range(date_from, date_to, default_days=1)
    data = await hisobot.report(db, kind, start, end, _filters(faculty, course, group, unit_kind, unit, q), criterion,
                                limit=None)
    content = build_workbook(data)
    label = next(c["label"] for c in data["criteria"] if c["key"] == data["criterion"])
    name = f"hisobot-{'talabalar' if kind == 'talaba' else 'xodimlar'}-{data['criterion']}-{start}_{end}.xlsx"
    return Response(
        content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}", "X-Report-Title": quote(label)},
    )


TabelKindQuery = Annotated[Literal["talaba", "xodim"], Query()]


@router.get("/tabel")
async def tabel(
    db: DbDep, _: ReadDep, kind: TabelKindQuery, oy: str,
    faculty: str | None = None, course: Annotated[int | None, Query(ge=1, le=7)] = None,
    group: str | None = None, unit_kind: str | None = None, unit: str | None = None, q: str | None = None,
) -> dict:
    """Oylik tabel: qatorlar — odamlar, ustunlar — oyning kunlari."""
    return await tabel_svc.build(db, kind, oy, _filters(faculty, course, group, unit_kind, unit, q))


@router.get("/tabel.xlsx")
async def tabel_export(
    db: DbDep, _: ReadDep, kind: TabelKindQuery, oy: str,
    faculty: str | None = None, course: Annotated[int | None, Query(ge=1, le=7)] = None,
    group: str | None = None, unit_kind: str | None = None, unit: str | None = None, q: str | None = None,
) -> Response:
    data = await tabel_svc.build(db, kind, oy, _filters(faculty, course, group, unit_kind, unit, q))
    name = f"tabel-{'talabalar' if kind == 'talaba' else 'xodimlar'}-{data['month']}.xlsx"
    return Response(
        build_tabel_workbook(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}",
                 "X-Report-Title": quote(data["title"])},
    )


# Chop etiladigan varaq: sarlavha bloki, jadval, izoh, imzo joyi.
_WEEKEND_FILL = PatternFill("solid", fgColor="EDEDED")
_HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
_THIN = Side(style="thin", color="999999")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_TOTAL_COLUMNS = [("present", "Keldi"), ("late", "Kech keldi"), ("absent", "Kelmadi"),
                  ("unknown", "Ma'lumot yo'q"), ("workDays", "Ish kunlari")]


def build_tabel_workbook(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tabel"
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")

    ws.append([settings.org_name])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([data["title"]])
    ws["A2"].font = Font(bold=True, size=14)
    ws.append([data["scope"]])
    ws.append([data["monthLabel"]])
    ws.append([])

    days = data["days"]
    # 1-4 sarlavha bloki, 5 bo'sh, 6 — kun raqamlari, 7 — hafta kunlari.
    head = 6
    first_day_col = 4
    ws.cell(head, 1, "№").font = bold
    ws.cell(head, 2, "F.I.Sh.").font = bold
    ws.cell(head, 3, "Guruh / bo'linma").font = bold
    for i, day in enumerate(days):
        col = first_day_col + i
        top = ws.cell(head, col, day["day"])
        sub = ws.cell(head + 1, col, day["weekday"])
        for cell in (top, sub):
            cell.font = bold
            cell.alignment = center
            cell.border = _BOX
            cell.fill = _HEADER_FILL if day["isWorkDay"] else _WEEKEND_FILL
        ws.column_dimensions[get_column_letter(col)].width = 4
    total_col = first_day_col + len(days)
    for i, (_key, label) in enumerate(_TOTAL_COLUMNS):
        cell = ws.cell(head, total_col + i, label)
        cell.font = bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BOX
        ws.column_dimensions[get_column_letter(total_col + i)].width = 9
    for col in (1, 2, 3):
        for row in (head, head + 1):
            ws.cell(row, col).fill = _HEADER_FILL
            ws.cell(row, col).border = _BOX
        ws.merge_cells(start_row=head, start_column=col, end_row=head + 1, end_column=col)
    for i in range(len(_TOTAL_COLUMNS)):
        ws.merge_cells(start_row=head, start_column=total_col + i, end_row=head + 1, end_column=total_col + i)

    row = head + 2
    for n, person in enumerate(data["people"], start=1):
        ws.cell(row, 1, n).alignment = center
        ws.cell(row, 2, person["fullName"])
        ws.cell(row, 3, person["group"])
        for i, (cell_data, day) in enumerate(zip(person["cells"], days)):
            cell = ws.cell(row, first_day_col + i, cell_data["mark"])
            cell.alignment = center
            cell.border = _BOX
            if not day["isWorkDay"]:
                cell.fill = _WEEKEND_FILL
        for i, (key, _label) in enumerate(_TOTAL_COLUMNS):
            cell = ws.cell(row, total_col + i, person["totals"][key])
            cell.alignment = center
            cell.border = _BOX
        for col in (1, 2, 3):
            ws.cell(row, col).border = _BOX
        row += 1

    ws.freeze_panes = f"C{head + 2}"  # birinchi ikki ustun va sarlavha qatori
    ws.print_title_rows = f"{head}:{head + 1}"
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 22

    row += 1
    ws.cell(row, 1, "Belgilar").font = bold
    row += 1
    for item in data["legend"]:
        ws.cell(row, 1, item["mark"]).alignment = center
        ws.cell(row, 2, item["label"])
        row += 1
    if data.get("note"):
        row += 1
        ws.cell(row, 1, "Izoh")
        ws.cell(row, 2, data["note"])
        row += 1
    row += 2
    ws.cell(row, 1, "Tabelni to'ldirgan: ______________________  /______________________/")
    row += 2
    ws.cell(row, 1, "Tasdiqlayman: ______________________  /______________________/")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_workbook(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hisobot"
    bold = Font(bold=True)
    label = next(c["label"] for c in data["criteria"] if c["key"] == data["criterion"])
    body = data["report"]
    # Ekrandagi bilan bir xil: sarlavha, tanlov, davr, keyin javob gapi.
    ws.append([f"{'Talabalar' if data['kind'] == 'talaba' else 'Xodimlar'} — {label}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Tanlov", data.get("scope", "")])
    ws.append(["Davr", f"{data['period']['from']} — {data['period']['to']}"])
    ws.append(["Ro'yxatdagi odamlar", data["population"]["total"]])
    ws.append([])
    for line in body.get("summary") or []:
        ws.append([line])
    ws.append([])
    for tile in body["tiles"]:
        ws.append([tile["label"], f"{tile['value']} {tile['unit']}".strip(), tile["hint"] or ""])
    if body.get("note"):
        ws.append(["Izoh", body["note"]])
    ws.append([])

    if body["breakdown"] and body["breakdown"]["rows"]:
        ws.append([body["breakdown"]["title"]])
        ws.cell(ws.max_row, 1).font = bold
        ws.append(["Nomi", f"Qiymat ({body['breakdown']['unit']})", "Izoh", "Odamlar soni"])
        for row in body["breakdown"]["rows"]:
            ws.append([row["name"], row["value"], row["detail"] or "", row["headcount"]])
        ws.append([])

    if body["columns"]:
        people = wb.create_sheet("Odamlar")
        people.append([body.get("people_title") or "Odamlar"])
        people.cell(1, 1).font = Font(bold=True, size=12)
        people.append([body.get("people_hint") or ""])
        people.append([])
        # Ustunlar ekrandagi jadval bilan bir xil tartibda va bir xil nom bilan.
        header = ["№", "F.I.Sh.", "Guruh yoki bo'linma"] + [
            f"{c['label']}, {c['unit']}" if c["unit"] else c["label"] for c in body["columns"]]
        people.append(header)
        for cell in people[people.max_row]:
            cell.font = bold
        for i, row in enumerate(body["people"], start=1):
            people.append([i, row["full_name"], row["unit"]] + [row["values"].get(c["key"]) for c in body["columns"]])
        people.column_dimensions["B"].width = 36
        people.column_dimensions["C"].width = 36
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 30
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
