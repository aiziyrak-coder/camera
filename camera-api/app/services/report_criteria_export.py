"""Hisobotlar sahifasining Excel nusxasi — GET /api/reports/criteria.xlsx.

Sahifada tanlangan filtr bo'yicha (bo'lim, davr va — ro'yxat ochilgan
bo'lsa — kriteriya, guruh, qidiruv) ikki xil fayl:

  * qisqa — faqat raqamlar: har kriteriya va uning har bir ko'rsatkichi
    ("Keldi — 312", "Kelmadi — 41" ...);
  * to'liq — xuddi shu raqamlar va har raqam ortidagi odamlar ism-familiyasi
    bilan: davomatda har kun-yozuv (sana, holat, kelgan vaqti va uni birinchi
    ko'rgan kamera), boshqa kriteriyalarda odamlar ro'yxati, signal
    kriteriyalarida har signal (vaqt, kamera, kim, holat).

Raqamlar sahifadagi bilan bir xil manbadan (app/services/report_criteria.py)
olinadi — fayl va ekran bir-biridan farq qilmasligi uchun.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import AttendanceRecord, Camera, Event, Faculty, PresenceVisit, StudentStaff
from app.schemas.report_criteria import ReportCriterionOut
from app.services.report_criteria import (
    Period,
    build_criteria,
    decorate_people,
    people_query,
)
from app.timezone import to_local

HEADER_FILL = PatternFill("solid", fgColor="E0E7FF")
HEADER_FONT = Font(bold=True, color="1E1B4B")
TITLE_FONT = Font(bold=True, size=14, color="1E1B4B")
SECTION_FONT = Font(bold=True, size=11, color="334155")

STATUS_LABELS = {"keldi": "Keldi", "kech_keldi": "Kechikdi", "kelmadi": "Kelmadi", "dam_olish": "Dam olish"}
EVENT_STATUS_LABELS = {"yangi": "Ko'rilmagan", "tasdiqlangan": "Tasdiqlangan", "rad_etilgan": "Rad etilgan"}
# Excel varaq nomida taqiqlangan belgilar.
_SHEET_FORBIDDEN = str.maketrans({ch: " " for ch in "[]:*?/\\"})


def _header(ws: Worksheet, values: list[str]) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _sheet(wb: Workbook, title: str) -> Worksheet:
    name = title.translate(_SHEET_FORBIDDEN).strip()[:31] or "Varaq"
    base, n = name, 2
    while name in wb.sheetnames:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    return wb.create_sheet(name)


def _hm(value) -> str:
    return value.strftime("%H:%M") if value else ""


def _summary(ws: Worksheet, report, criteria: list[ReportCriterionOut], variant_label: str) -> None:
    ws.title = "Xulosa"
    ws.append([f"Hisobot — {report.population_label}"])
    ws.cell(row=1, column=1).font = TITLE_FONT
    period = report.period
    span = period.start if period.start == period.end else f"{period.start} — {period.end}"
    ws.append([f"Davr: {period.label} ({span})"])
    ws.append([f"Jami: {report.people_total} ta, yuzi ro'yxatdan o'tgan: {report.enrolled_total} ta"])
    ws.append([f"Fayl: {variant_label}"])
    ws.append([])
    _header(ws, ["Kriteriya", "Ko'rsatkich", "Soni", "Izoh"])
    for criterion in criteria:
        for index, bucket in enumerate(criterion.buckets):
            ws.append(
                [
                    criterion.title if index == 0 else "",
                    bucket.label,
                    bucket.count,
                    (criterion.note or "") if index == 0 else "",
                ]
            )
        ws.append(["", f"Jami ({criterion.unit})" if criterion.unit else "Jami", criterion.total, ""])
    _widths(ws, [30, 22, 10, 60])


async def _first_cameras(db: AsyncSession, person_ids: list, period: Period) -> dict:
    """(odam, mahalliy sana) -> shu kuni uni birinchi ko'rgan kamera — isbot."""
    if not person_ids:
        return {}
    start, end = period.bounds_utc()
    first: dict = {}
    rows = (
        await db.execute(
            select(PresenceVisit.student_staff_id, PresenceVisit.first_seen_at, Camera.name)
            .join(Camera, Camera.id == PresenceVisit.camera_id)
            .where(PresenceVisit.student_staff_id.in_(person_ids))
            .where(PresenceVisit.first_seen_at.between(start, end))
            .order_by(PresenceVisit.first_seen_at)
        )
    ).all()
    for person_id, seen_at, camera_name in rows:
        first.setdefault((person_id, to_local(seen_at).date()), camera_name)
    return first


async def _attendance_sheet(
    wb: Workbook,
    db: AsyncSession,
    population: str,
    period: Period,
    criterion: ReportCriterionOut,
    bucket: str,
    search: str | None,
) -> None:
    stmt = (
        select(AttendanceRecord, StudentStaff)
        .join(StudentStaff, StudentStaff.id == AttendanceRecord.student_staff_id)
        .where(StudentStaff.type == population)
        .where(AttendanceRecord.date.between(period.start, period.end))
    )
    if bucket:
        stmt = stmt.where(AttendanceRecord.status == bucket)
    if search:
        stmt = stmt.where(StudentStaff.full_name.ilike(f"%{search}%"))
    stmt = stmt.order_by(AttendanceRecord.date, StudentStaff.full_name)
    rows = (await db.execute(stmt)).all()

    faculty_names: dict = {}
    if rows:
        faculty_ids = {person.faculty_id for _, person in rows if person.faculty_id}
        if faculty_ids:
            faculty_names = dict(
                (await db.execute(select(Faculty.id, Faculty.name).where(Faculty.id.in_(faculty_ids)))).all()
            )
    cameras = await _first_cameras(db, list({person.id for _, person in rows}), period)

    ws = _sheet(wb, criterion.title)
    _header(
        ws,
        ["№", "Sana", "F.I.Sh", "Fakultet", "Lavozim / guruh", "Holat", "Kelgan vaqti", "Ketgan vaqti", "Birinchi ko'rgan kamera"],
    )
    for index, (record, person) in enumerate(rows, start=1):
        ws.append(
            [
                index,
                record.date.isoformat(),
                person.full_name,
                faculty_names.get(person.faculty_id, ""),
                person.group_or_position,
                STATUS_LABELS.get(record.status, record.status),
                _hm(record.check_in),
                _hm(record.check_out),
                cameras.get((person.id, record.date), ""),
            ]
        )
    if not rows:
        ws.append(["", "Bu filtr bo'yicha yozuv yo'q"])
    _widths(ws, [6, 12, 34, 26, 26, 12, 13, 13, 30])
    ws.freeze_panes = "A2"


async def _people_sheet(
    wb: Workbook,
    db: AsyncSession,
    population: str,
    period: Period,
    criterion: ReportCriterionOut,
    bucket: str,
    search: str | None,
) -> None:
    ws = _sheet(wb, criterion.title)
    _header(
        ws,
        [
            "№",
            "F.I.Sh",
            "Fakultet",
            "Lavozim / guruh",
            "Ko'rsatkich",
            "Kelgan kunlar",
            "Kechikkan",
            "Kelmagan",
            "Birinchi kelish",
            "Kamerada ko'rinish",
            "Kameralar soni",
            "Oxirgi ko'rgan kamera",
            "Oxirgi ko'rilgan",
        ],
    )
    number = 0
    buckets = [b for b in criterion.buckets if not bucket or b.key == bucket]
    for item in buckets:
        stmt = await people_query(db, population, criterion.key, item.key, period)
        if search:
            stmt = stmt.where(StudentStaff.full_name.ilike(f"%{search}%"))
        people = list((await db.execute(stmt)).scalars().all())
        for row in await decorate_people(db, people, period):
            number += 1
            last_seen = ""
            if row.last_seen_at:
                last_seen = to_local(datetime.fromisoformat(row.last_seen_at)).strftime("%Y-%m-%d %H:%M")
            ws.append(
                [
                    number,
                    row.full_name,
                    row.faculty,
                    row.unit,
                    item.label,
                    row.present_days,
                    row.late_days,
                    row.absent_days,
                    row.first_check_in or "",
                    row.visits,
                    row.cameras,
                    row.last_seen_camera or "",
                    last_seen,
                ]
            )
    if number == 0:
        ws.append(["", "Bu filtr bo'yicha odam yo'q"])
    _widths(ws, [6, 34, 26, 26, 16, 13, 11, 11, 14, 14, 12, 30, 18])
    ws.freeze_panes = "A2"


async def _events_sheet(
    wb: Workbook, db: AsyncSession, period: Period, criterion: ReportCriterionOut, bucket: str
) -> None:
    start, end = period.bounds_utc()
    stmt = (
        select(Event)
        .where(Event.is_trial.is_(False))
        .where(Event.module_code.in_(criterion.module_codes))
        .where(Event.occurred_at.between(start, end))
        .order_by(Event.occurred_at)
    )
    if bucket:
        stmt = stmt.where(Event.status == bucket)
    events = list((await db.execute(stmt)).scalars().all())
    ws = _sheet(wb, criterion.title)
    _header(ws, ["№", "Vaqt", "Mezon", "Kamera", "Bino", "Kim", "Holat", "Ishonch, %"])
    for index, event in enumerate(events, start=1):
        ws.append(
            [
                index,
                to_local(event.occurred_at).strftime("%Y-%m-%d %H:%M:%S"),
                event.module_name,
                event.camera_name,
                event.building,
                event.person_name or "",
                EVENT_STATUS_LABELS.get(event.status, event.status),
                event.confidence,
            ]
        )
    if not events:
        ws.append(["", "Bu davrda signal yo'q"])
    _widths(ws, [6, 20, 30, 30, 22, 30, 14, 11])
    ws.freeze_panes = "A2"


async def build_criteria_workbook(
    db: AsyncSession,
    population: str,
    period: Period,
    *,
    variant: str,
    criterion_key: str | None = None,
    bucket: str = "",
    search: str | None = None,
) -> bytes:
    """variant: "qisqa" (faqat raqamlar) yoki "toliq" (raqamlar + ism-familiyalar)."""
    report = await build_criteria(db, population, period)
    criteria = report.criteria
    if criterion_key:
        criteria = [c for c in criteria if c.key == criterion_key] or criteria
    search = (search or "").strip() or None
    selected_bucket = bucket if criterion_key else ""

    wb = Workbook()
    _summary(wb.active, report, criteria, "to'liq (ism-familiyalar bilan)" if variant == "toliq" else "qisqa (raqamlar)")

    if variant == "toliq":
        for criterion in criteria:
            if criterion.key == "davomat":
                await _attendance_sheet(wb, db, population, period, criterion, selected_bucket, search)
            elif criterion.detail == "people":
                await _people_sheet(wb, db, population, period, criterion, selected_bucket, search)
            elif criterion.detail == "events":
                await _events_sheet(wb, db, period, criterion, selected_bucket)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
