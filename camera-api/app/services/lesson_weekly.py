"""Haftalik dars jadvali — namuna fayl va uni semestrga yoyish.

Nega haftalik? Kafedra mudiri jadvalni HAFTA bo'yicha biladi: "DI-2301,
seshanba, 08:30, 214-xona". Sana bo'yicha so'ralsa, bitta semestr uchun
o'sha odam mingdan ortiq qator yozishi kerak bo'lardi — amalda hech kim
to'ldirmaydi. Shuning uchun fayl haftalik, tizim esa uni berilgan
oraliqdagi haqiqiy sanalarga yoyadi va keyin mavjud import quvuriga
(app/services/lesson_import.py) uzatadi: xona -> kamera, o'qituvchi
ismi -> xodim, takrorni o'tkazib yuborish — hammasi o'sha yerda.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date as date_type, timedelta

WEEKDAY_NAMES: tuple[str, ...] = (
    "Dushanba",
    "Seshanba",
    "Chorshanba",
    "Payshanba",
    "Juma",
    "Shanba",
    "Yakshanba",
)

# Qabul qilinadigan yozilishlar -> ISO hafta kuni (1 = dushanba).
_WEEKDAY_ALIASES: dict[str, int] = {}
for _index, _name in enumerate(WEEKDAY_NAMES, start=1):
    for _form in (_name, _name[:2], _name[:3], str(_index)):
        _WEEKDAY_ALIASES[_form.lower()] = _index
# Ko'p uchraydigan boshqa yozilishlar.
_WEEKDAY_ALIASES.update(
    {
        "du": 1, "dush": 1, "понедельник": 1, "пн": 1,
        "se": 2, "sesh": 2, "вторник": 2, "вт": 2,
        "ch": 3, "chor": 3, "среда": 3, "ср": 3,
        "pa": 4, "pay": 4, "четверг": 4, "чт": 4,
        "ju": 5, "жума": 5, "пятница": 5, "пт": 5,
        "sh": 6, "shan": 6, "суббота": 6, "сб": 6,
        "ya": 7, "yak": 7, "воскресенье": 7, "вс": 7,
    }
)

WEEKDAY_COLUMN_ALIASES = ("hafta_kuni", "hafta", "kun", "weekday", "day", "dars_kuni")

#: Bitta yuklashda tuziladigan darslar chegarasi. 285 guruh × 6 kun ×
#: 6 para × 18 hafta juda katta son beradi; chegarasiz bitta noto'g'ri
#: oraliq bazani ko'mib tashlaydi.
MAX_GENERATED = 40_000


def parse_weekday(value: object) -> int | None:
    """"Seshanba", "se", "2" -> 2. Tushunilmasa None."""
    text = str(value or "").strip().lower()
    if not text:
        return None
    for ch in "'‘’`ʻʼ":
        text = text.replace(ch, "")
    text = re.sub(r"[\s.\-]+", "", text)
    return _WEEKDAY_ALIASES.get(text)


@dataclass
class WeeklyExpansion:
    """Yoyish natijasi: sanali qatorlar va nima bo'lganining hisobi."""

    rows: list[dict[str, object]]
    header: list[str]
    lessons: int
    weeks: int
    bad_weekday: list[int]
    """Hafta kuni tushunilmagan qatorlar (fayldagi raqami)."""
    truncated: bool


def weekday_dates(weekday: int, start: date_type, end: date_type) -> list[date_type]:
    """Oraliqdagi barcha shu hafta kuni sanalari."""
    if start > end:
        return []
    # Birinchi mos kun: start dan boshlab oldinga siljiymiz.
    offset = (weekday - start.isoweekday()) % 7
    current = start + timedelta(days=offset)
    out: list[date_type] = []
    while current <= end:
        out.append(current)
        current += timedelta(days=7)
    return out


def expand_weekly(
    header: list[str],
    rows: list[dict[str, object]],
    start: date_type,
    end: date_type,
    *,
    weekday_column: str,
) -> WeeklyExpansion:
    """Haftalik qatorlarni oraliqdagi sanali qatorlarga aylantiradi.

    Har bir chiqish qatori kirish qatorining nusxasi bo'lib, unga `sana`
    ustuni qo'shiladi — keyin uni odatdagi import o'qiydi.
    """
    out: list[dict[str, object]] = []
    bad: list[int] = []
    weeks = 0
    truncated = False

    for row_num, row in enumerate(rows, start=2):
        if not any(str(value or "").strip() for value in row.values()):
            continue
        weekday = parse_weekday(row.get(weekday_column))
        if weekday is None:
            bad.append(row_num)
            continue
        dates = weekday_dates(weekday, start, end)
        weeks = max(weeks, len(dates))
        for lesson_date in dates:
            if len(out) >= MAX_GENERATED:
                truncated = True
                break
            copy = dict(row)
            copy.pop(weekday_column, None)
            copy["sana"] = lesson_date.isoformat()
            out.append(copy)
        if truncated:
            break

    new_header = [name for name in header if name != weekday_column]
    if "sana" not in new_header:
        new_header.append("sana")
    return WeeklyExpansion(
        rows=out,
        header=new_header,
        lessons=len(out),
        weeks=weeks,
        bad_weekday=bad,
        truncated=truncated,
    )


# ─────────────────────────────────────────────── Namuna (.xlsx)

TEMPLATE_HEADERS = ("Guruh", "Hafta kuni", "Boshlanish", "Xona", "Fan", "O'qituvchi")

_INSTRUCTIONS = [
    ("1", "Har bir HAFTALIK dars uchun bitta qator to'ldiring."),
    ("2", "Guruh va Xona — ro'yxatdan tanlanadi (katakni bosing)."),
    ("3", "Boshlanish — 08:30 ko'rinishida."),
    ("4", "Sana yozilmaydi: tizim jadvalni butun semestrga yoyadi."),
    ("5", "Fan va O'qituvchi bo'sh qolsa ham bo'ladi."),
    ("6", "Faylni o'zgartirmang: ustun nomlari shu holda qolsin."),
]


def build_weekly_template(
    groups: list[str],
    rooms: list[tuple[str, str]],
    *,
    lesson_minutes: int,
) -> bytes:
    """Kafedraga yuboriladigan namuna: to'ldiriladigan varaq + ro'yxatlar.

    `rooms` — (xona kodi, kamera nomi). Faqat KAMERASI BOR xonalar
    beriladi: kamerasiz xonadagi darsni tizim baribir tekshira olmaydi,
    shuning uchun uni ro'yxatga qo'shib umid uyg'otmaymiz.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Jadval"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DCE6F1")

    for index, title in enumerate(TEMPLATE_HEADERS, start=1):
        cell = sheet.cell(1, index, title)
        cell.font = bold
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")
    widths = (16, 14, 14, 12, 28, 30)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"

    # Namuna qator — o'chirib, o'rniga o'zinikini yozish mumkin.
    example = (groups[0] if groups else "DI-2301", "Seshanba", "08:30", rooms[0][0] if rooms else "214", "Anatomiya", "")
    for index, value in enumerate(example, start=1):
        cell = sheet.cell(2, index, value)
        cell.font = Font(italic=True, color="808080")

    ref = wb.create_sheet("Guruhlar")
    ref.cell(1, 1, "Guruhlar").font = bold
    for index, name in enumerate(groups, start=2):
        ref.cell(index, 1, name)
    ref.column_dimensions["A"].width = 22

    rooms_sheet = wb.create_sheet("Xonalar")
    rooms_sheet.cell(1, 1, "Xona").font = bold
    rooms_sheet.cell(1, 2, "Kamera").font = bold
    for index, (code, camera) in enumerate(rooms, start=2):
        rooms_sheet.cell(index, 1, code)
        rooms_sheet.cell(index, 2, camera)
    rooms_sheet.column_dimensions["A"].width = 12
    rooms_sheet.column_dimensions["B"].width = 32

    guide = wb.create_sheet("Yo'riqnoma")
    guide.cell(1, 1, "Dars jadvali — namuna").font = Font(bold=True, size=13)
    for index, (number, text) in enumerate(_INSTRUCTIONS, start=3):
        guide.cell(index, 1, number).font = bold
        guide.cell(index, 2, text)
    tail = 3 + len(_INSTRUCTIONS) + 1
    guide.cell(tail, 2, f"Bitta dars {lesson_minutes} daqiqa deb hisoblanadi.")
    if not rooms:
        guide.cell(tail + 1, 2, "Diqqat: hozircha birorta xonaga kamera biriktirilmagan.")
    guide.column_dimensions["A"].width = 4
    guide.column_dimensions["B"].width = 70

    # Ro'yxatdan tanlash — noto'g'ri guruh yoki xona yozilishining oldini oladi.
    last = 500
    if groups:
        dv = DataValidation(type="list", formula1=f"=Guruhlar!$A$2:$A${len(groups) + 1}", allow_blank=True)
        dv.error = "Guruhni ro'yxatdan tanlang"
        dv.errorTitle = "Noto'g'ri guruh"
        sheet.add_data_validation(dv)
        dv.add(f"A2:A{last}")
    weekdays = DataValidation(type="list", formula1=f'"{",".join(WEEKDAY_NAMES)}"', allow_blank=True)
    weekdays.error = "Hafta kunini ro'yxatdan tanlang"
    weekdays.errorTitle = "Noto'g'ri kun"
    sheet.add_data_validation(weekdays)
    weekdays.add(f"B2:B{last}")
    if rooms:
        dv_rooms = DataValidation(type="list", formula1=f"=Xonalar!$A$2:$A${len(rooms) + 1}", allow_blank=True)
        dv_rooms.error = "Kamerasi bor xonani tanlang"
        dv_rooms.errorTitle = "Noto'g'ri xona"
        sheet.add_data_validation(dv_rooms)
        dv_rooms.add(f"D2:D{last}")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
