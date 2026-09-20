"""Dars jadvalini ommaviy import qilish — POST /api/lesson-sessions/import.

Dars mezonlari (#8, #19, #21, #22, #26 va dars davomati) jadvalsiz umuman
ishlamaydi: 2026-09-18 da tizimda 0 ta dars bor edi. Eski import esa
kamera va o'qituvchini faqat ichki UUID bilan qabul qilardi — o'quv
bo'limidagi haqiqiy jadvalda esa xona raqami va o'qituvchining ismi
bo'ladi. Endi:

  * CSV yoki Excel (.xlsx), o'zbekcha sarlavhalar ham (sana, guruh, fan,
    xona, o'qituvchi, boshlanish);
  * xona -> kamera: Camera.room_code orqali ("211-xona" = "211",
    app/services/camera_roles.normalize_room_code);
  * o'qituvchi ismi -> xodim: app/services/name_matching (Saloxiddin/
    Salohiddin, kirill yozuvi va h.k.);
  * fakultet ustuni bo'lmasa — guruh talabalarining fakultetidan;
  * sana: 2026-09-21 yoki 21.09.2026; vaqt: to'liq ISO, "08:30" yoki
    "08:30-09:50" (boshlanishi olinadi);
  * `apply=False` — faqat oldindan ko'rish: nima qo'shilishi, qaysi xona
    yoki o'qituvchi topilmagani. Hech narsa yozilmaydi.

Topilmagan xona yoki o'qituvchi darsni to'xtatmaydi: dars kamerasiz/
o'qituvchisiz qo'shiladi va ogohlantirishda ko'rsatiladi (keyin kamera
xona raqami qo'yilgach qayta import qilinsa — takror qatorlar o'tkazib
yuboriladi).
"""

from __future__ import annotations

import csv
import io
import logging
import re
import uuid
from collections import Counter
from dataclasses import dataclass, field
from datetime import date as date_type, datetime, time as time_type

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Camera, Faculty, LessonSession, StudentStaff
from app.schemas.lesson_session import (
    LessonImportPreviewRowOut,
    LessonSessionImportErrorOut,
    LessonSessionImportResultOut,
)
from app.services.camera_roles import normalize_room_code
from app.services.name_matching import name_key, same_person_name
from app.timezone import INSTITUTE_TZ

logger = logging.getLogger("app.lesson_import")

# Kanonik ustun -> fayldagi mumkin bo'lgan nomlar (normallashgan).
_ALIASES: dict[str, tuple[str, ...]] = {
    "date": ("date", "sana", "kun"),
    "group": ("group", "guruh", "group_name"),
    "faculty": ("faculty", "fakultet"),
    "subject": ("subject", "fan", "fan_nomi"),
    "teacher_id": ("teacher_id",),
    "teacher": ("teacher", "teacher_name", "oqituvchi", "oqituvchi_fio", "fio"),
    "camera_id": ("camera_id",),
    "room": ("room", "room_code", "xona", "auditoriya", "xona_raqami"),
    "scheduled_start_time": (
        "scheduled_start_time",
        "start",
        "start_time",
        "boshlanish",
        "boshlanish_vaqti",
        "vaqt",
        "dars_vaqti",
    ),
}
_REQUIRED = ("date", "group", "subject")
PREVIEW_ROWS = 200

LessonKey = tuple[date_type, str, str, datetime | None]


def parse_scheduled_start_time(value: str | None) -> datetime | None:
    """ISO 8601 schedule time — naive values use institute-local TZ."""
    if value is None:
        return None
    parsed = datetime.fromisoformat(value)
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=INSTITUTE_TZ)
    return parsed


def _normalize_header(name: object) -> str:
    text = str(name or "").strip().lower()
    for ch in "'‘’`ʻʼ":
        text = text.replace(ch, "")
    return re.sub(r"[\s\-]+", "_", text)


_DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d.%m.%y")


def _parse_date(value: object) -> date_type | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date_type):
        return value
    text = str(value or "").strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


_CLOCK = re.compile(r"^\s*(\d{1,2})[:.](\d{2})")


def _parse_start(value: object, lesson_date: date_type) -> datetime | None:
    """To'liq sana-vaqt, "08:30", "8.30" yoki "08:30-09:50". Bo'sh — None.
    Tushunilmagan qiymat ValueError."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=INSTITUTE_TZ)
    if isinstance(value, time_type):
        return datetime.combine(lesson_date, value, tzinfo=INSTITUTE_TZ)
    text = str(value).strip()
    if not text:
        return None
    match = _CLOCK.match(text)
    if match and len(text) <= 13:  # "08:30" yoki "08:30-09:50"
        hour, minute = int(match.group(1)), int(match.group(2))
        return datetime.combine(lesson_date, time_type(hour, minute), tzinfo=INSTITUTE_TZ)
    return parse_scheduled_start_time(text)


def _read_rows(raw: bytes, filename: str) -> tuple[list[str], list[dict[str, object]]]:
    """(normallashgan sarlavhalar, qatorlar) — CSV yoki birinchi Excel varag'i."""
    if filename.lower().endswith(".xlsx") or raw[:2] == b"PK":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows_iter = sheet.iter_rows(values_only=True)
        header: list[str] = []
        for values in rows_iter:
            if any(cell not in (None, "") for cell in values):
                header = [_normalize_header(cell) for cell in values]
                break
        rows = [
            {header[i]: cell for i, cell in enumerate(values) if i < len(header) and header[i]}
            for values in rows_iter
        ]
        workbook.close()
        return header, rows

    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    first_line = text.split("\n", 1)[0]
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    header_raw = next(reader, None)
    if header_raw is None:
        return [], []
    header = [_normalize_header(cell) for cell in header_raw]
    rows = [{header[i]: cell for i, cell in enumerate(values) if i < len(header) and header[i]} for values in reader]
    return header, rows


def _column_map(header: list[str]) -> dict[str, str]:
    present = set(header)
    mapping: dict[str, str] = {}
    for canonical, aliases in _ALIASES.items():
        for alias in aliases:
            if alias in present:
                mapping[canonical] = alias
                break
    return mapping


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel "211" ni 211.0 qilib beradi
    return str(value).strip()


def _parse_uuid(value: str | None) -> uuid.UUID | None:
    if not value:
        return None
    try:
        return uuid.UUID(value)
    except ValueError:
        return None


@dataclass
class _Lookups:
    faculty_names: set[str]
    existing_keys: set[LessonKey]
    staff: list[StudentStaff]
    staff_by_id: dict[uuid.UUID, StudentStaff]
    cameras_by_id: dict[uuid.UUID, Camera]
    cameras_by_room: dict[str, Camera]
    teacher_cache: dict[str, StudentStaff | None] = field(default_factory=dict)
    faculty_cache: dict[str, str | None] = field(default_factory=dict)


async def _load_lookups(db: AsyncSession) -> _Lookups:
    faculty_names = set((await db.execute(select(Faculty.name))).scalars().all())
    existing = await db.execute(
        select(LessonSession.date, LessonSession.group_name, LessonSession.subject, LessonSession.scheduled_start_time)
    )
    staff = list((await db.execute(select(StudentStaff).where(StudentStaff.type == "xodim"))).scalars().all())
    cameras = list((await db.execute(select(Camera).order_by(Camera.name))).scalars().all())
    by_room: dict[str, Camera] = {}
    for camera in cameras:
        # Bir xonada bir nechta kamera bo'lsa — nomi bo'yicha birinchisi
        # (tartib barqaror: qayta import o'sha kamerani tanlaydi).
        if camera.room_code and camera.room_code not in by_room:
            by_room[camera.room_code] = camera
    return _Lookups(
        faculty_names=faculty_names,
        existing_keys={tuple(row) for row in existing.all()},
        staff=staff,
        staff_by_id={person.id: person for person in staff},
        cameras_by_id={camera.id: camera for camera in cameras},
        cameras_by_room=by_room,
    )


def _find_teacher(lookups: _Lookups, name: str) -> StudentStaff | None:
    """Aniq ism, bo'lmasa name_matching qoidasi. Bir nechta xodimga mos
    kelsa — hech biri (noto'g'ri o'qituvchini kuzatgandan ko'ra bo'sh
    qoldirgan yaxshi)."""
    key = name_key(name)
    if key in lookups.teacher_cache:
        return lookups.teacher_cache[key]
    exact = [person for person in lookups.staff if name_key(person.full_name) == key]
    found = exact if exact else [person for person in lookups.staff if same_person_name(person.full_name, name)]
    teacher = found[0] if len(found) == 1 else None
    lookups.teacher_cache[key] = teacher
    return teacher


async def _group_faculty(db: AsyncSession, lookups: _Lookups, group: str) -> str | None:
    """Guruh talabalarining eng ko'p uchraydigan fakulteti."""
    if group in lookups.faculty_cache:
        return lookups.faculty_cache[group]
    from app.jobs.lesson_attendance import group_member_clause

    rows = await db.execute(
        select(Faculty.name)
        .join(StudentStaff, StudentStaff.faculty_id == Faculty.id)
        .where(StudentStaff.type == "talaba")
        .where(group_member_clause(group))
    )
    counts = Counter(rows.scalars().all())
    faculty = counts.most_common(1)[0][0] if counts else None
    lookups.faculty_cache[group] = faculty
    return faculty


async def import_lesson_sessions(
    db: AsyncSession,
    raw: bytes,
    *,
    filename: str = "",
    apply: bool = True,
    prepared: tuple[list[str], list[dict[str, object]]] | None = None,
) -> LessonSessionImportResultOut:
    """Jadvalni o'qiydi; `apply` bo'lsa darslarni sessiyaga qo'shadi (commit
    chaqiruvchida, audit yozuvi bilan birga).

    `prepared` — fayl allaqachon o'qilgan va qatorlar tayyorlangan bo'lsa
    (haftalik jadval sanalarga yoyilgandan keyin): bir xil quvurdan
    o'tsin, xona/o'qituvchi/takror mantig'i ikki joyda takrorlanmasin."""
    try:
        header, rows = prepared if prepared is not None else _read_rows(raw, filename)
    except Exception:
        logger.exception("lesson import: file could not be read")
        return LessonSessionImportResultOut(
            imported=0, skipped=0, errors=[LessonSessionImportErrorOut(row=0, message="Faylni o'qib bo'lmadi")]
        )
    if not header:
        return LessonSessionImportResultOut(
            imported=0, skipped=0, errors=[LessonSessionImportErrorOut(row=0, message="Fayl bo'sh")]
        )
    columns = _column_map(header)
    missing = [name for name in _REQUIRED if name not in columns]
    if missing:
        return LessonSessionImportResultOut(
            imported=0,
            skipped=0,
            errors=[LessonSessionImportErrorOut(row=0, message=f"Yetishmayotgan ustunlar: {', '.join(missing)}")],
        )

    lookups = await _load_lookups(db)
    pending: set[LessonKey] = set()
    result = LessonSessionImportResultOut(imported=0, skipped=0, errors=[], preview=not apply)
    unmatched_rooms: Counter[str] = Counter()
    unmatched_teachers: Counter[str] = Counter()

    def cell(row: dict[str, object], name: str) -> object:
        column = columns.get(name)
        return row.get(column) if column else None

    for row_num, row in enumerate(rows, start=2):
        if not any(_text(value) for value in row.values()):
            continue
        group = _text(cell(row, "group"))
        subject = _text(cell(row, "subject"))
        lesson_date = _parse_date(cell(row, "date"))
        if not group or not subject or cell(row, "date") in (None, ""):
            result.errors.append(LessonSessionImportErrorOut(row=row_num, message="Majburiy maydonlar to'ldirilmagan"))
            continue
        if lesson_date is None:
            result.errors.append(
                LessonSessionImportErrorOut(row=row_num, message=f"Noto'g'ri sana: {_text(cell(row, 'date'))}")
            )
            continue
        try:
            scheduled_start = _parse_start(cell(row, "scheduled_start_time"), lesson_date)
        except ValueError:
            result.errors.append(
                LessonSessionImportErrorOut(
                    row=row_num, message=f"Noto'g'ri vaqt: {_text(cell(row, 'scheduled_start_time'))}"
                )
            )
            continue

        faculty_name = _text(cell(row, "faculty"))
        if faculty_name and faculty_name not in lookups.faculty_names:
            result.errors.append(LessonSessionImportErrorOut(row=row_num, message=f"Fakultet topilmadi: {faculty_name}"))
            continue
        if not faculty_name:
            faculty_name = await _group_faculty(db, lookups, group) or ""
            if not faculty_name:
                result.errors.append(
                    LessonSessionImportErrorOut(
                        row=row_num, message=f"Fakultet ko'rsatilmagan va '{group}' guruhi talabalaridan aniqlanmadi"
                    )
                )
                continue

        teacher: StudentStaff | None = None
        teacher_text = _text(cell(row, "teacher"))
        teacher_id_raw = _text(cell(row, "teacher_id"))
        if teacher_id_raw:
            teacher_uuid = _parse_uuid(teacher_id_raw)
            if teacher_uuid is None:
                result.errors.append(LessonSessionImportErrorOut(row=row_num, message="O'qituvchi ID noto'g'ri"))
                continue
            teacher = lookups.staff_by_id.get(teacher_uuid)
            if teacher is None:
                result.errors.append(LessonSessionImportErrorOut(row=row_num, message="O'qituvchi topilmadi"))
                continue
        elif teacher_text:
            teacher = _find_teacher(lookups, teacher_text)
            if teacher is None:
                unmatched_teachers[teacher_text] += 1

        camera: Camera | None = None
        camera_id_raw = _text(cell(row, "camera_id"))
        room_text = _text(cell(row, "room"))
        if camera_id_raw:
            camera_uuid = _parse_uuid(camera_id_raw)
            if camera_uuid is None:
                result.errors.append(LessonSessionImportErrorOut(row=row_num, message="Kamera ID noto'g'ri"))
                continue
            camera = lookups.cameras_by_id.get(camera_uuid)
            if camera is None:
                result.errors.append(LessonSessionImportErrorOut(row=row_num, message="Kamera topilmadi"))
                continue
        elif room_text:
            camera = lookups.cameras_by_room.get(normalize_room_code(room_text) or "")
            if camera is None:
                unmatched_rooms[room_text] += 1

        lesson_key: LessonKey = (lesson_date, group, subject, scheduled_start)
        if lesson_key in lookups.existing_keys or lesson_key in pending:
            result.skipped += 1
            continue
        pending.add(lesson_key)
        result.imported += 1
        if camera is not None:
            result.with_camera += 1
        if teacher is not None:
            result.with_teacher += 1
        if len(result.rows) < PREVIEW_ROWS:
            result.rows.append(
                LessonImportPreviewRowOut(
                    row=row_num,
                    date=lesson_date.isoformat(),
                    start=scheduled_start.astimezone(INSTITUTE_TZ).strftime("%H:%M") if scheduled_start else None,
                    group=group,
                    subject=subject,
                    teacher=teacher.full_name if teacher else (teacher_text or None),
                    teacher_matched=teacher is not None,
                    room=room_text or None,
                    camera=camera.name if camera else None,
                )
            )
        if apply:
            db.add(
                LessonSession(
                    date=lesson_date,
                    group_name=group,
                    faculty=faculty_name,
                    subject=subject,
                    teacher=teacher.full_name if teacher else (teacher_text or "—"),
                    teacher_id=teacher.id if teacher else None,
                    camera_id=camera.id if camera else None,
                    scheduled_start_time=scheduled_start,
                    # O'lchovlar dars o'tganda yoziladi (lesson_quality_ai,
                    # teacher_punctuality_ai) — import ularni o'ylab topmaydi.
                    attention_score=0,
                    sleep_incidents=0,
                    teacher_activity_score=0,
                    teacher_on_time=None,
                )
            )

    result.unmatched_rooms = [name for name, _ in unmatched_rooms.most_common()]
    result.unmatched_teachers = [name for name, _ in unmatched_teachers.most_common()]
    return result


async def import_lesson_sessions_csv(db: AsyncSession, raw: bytes) -> LessonSessionImportResultOut:
    """Eski nom — to'g'ridan-to'g'ri import (oldindan ko'rishsiz)."""
    return await import_lesson_sessions(db, raw, apply=True)
