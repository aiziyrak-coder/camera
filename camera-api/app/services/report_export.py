"""Tahliliy hisobotning Excel (.xlsx) nusxasi — GET /api/reports/analytics.xlsx.

Sahifadagi bo'limlar alohida varaqlarda: Xulosa, Davomat, Xavfsizlik,
Darslar, Tizim. Grafik emas, jadval — Excel'ni ochgan odam raqamlarni
o'zi saralab/filtrlab ishlatadi. Uslub app/services/staff_export.py ga mos.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.report import AttendancePopulationOut, ReportAnalyticsOut

HEADER_FILL = PatternFill("solid", fgColor="E0E7FF")
HEADER_FONT = Font(bold=True, color="1E1B4B")
TITLE_FONT = Font(bold=True, size=14, color="1E1B4B")
SECTION_FONT = Font(bold=True, size=11, color="334155")
LEVEL_LABELS = {"critical": "Shoshilinch", "warning": "Diqqat", "info": "Ma'lumot", "ok": "Yaxshi"}


def _header(ws: Worksheet, values: list[str]) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _section(ws: Worksheet, title: str) -> None:
    if ws.max_row > 1:
        ws.append([])
    ws.append([title])
    ws.cell(row=ws.max_row, column=1).font = SECTION_FONT


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _pct(value: float | None) -> float | str:
    return "—" if value is None else value


def _attendance_rows(ws: Worksheet, population: AttendancePopulationOut) -> None:
    _section(ws, f"{population.label}: kunlar bo'yicha")
    _header(ws, ["Sana", "Keldi", "Kech keldi", "Kelmadi", "Davomat, %"])
    for day in population.by_day:
        ws.append([day.date, day.keldi, day.kech_keldi, day.kelmadi, _pct(day.rate)])
    if population.by_faculty:
        _section(ws, f"{population.label}: fakultetlar bo'yicha")
        _header(ws, ["Fakultet", "Yozuvlar", "Kelgan", "Kech", "Davomat, %"])
        for row in population.by_faculty:
            ws.append([row.name, row.total, row.present, row.late, _pct(row.rate)])
    for warning in population.reliability.warnings:
        ws.append([f"Ogohlantirish: {warning}"])


def build_analytics_workbook(a: ReportAnalyticsOut) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Xulosa"
    ws.append(["Farg'ona JSSTI — Situatsion Markaz: tahliliy hisobot"])
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([f"Davr: {a.period.label}", f"Solishtirish: {a.previous_period.label}", f"Tayyorlandi: {a.generated_at}"])
    _section(ws, "Asosiy ko'rsatkichlar")
    _header(ws, ["Ko'rsatkich", "Joriy davr", "Oldingi davr", "O'zgarish", "Izoh"])
    for kpi in a.kpis:
        ws.append([kpi.label, kpi.display, kpi.previous_display or "—", kpi.delta_display or "—", kpi.note or ""])
    _section(ws, "Asosiy xulosalar")
    _header(ws, ["Daraja", "Xulosa", "Tafsilot"])
    for insight in a.insights:
        ws.append([LEVEL_LABELS[insight.level], insight.title, insight.text])
    _widths(ws, [34, 36, 70, 16, 40])

    ws = wb.create_sheet("Davomat")
    _attendance_rows(ws, a.attendance.staff)
    _attendance_rows(ws, a.attendance.students)
    _widths(ws, [40, 12, 12, 12, 14])

    ws = wb.create_sheet("Xavfsizlik")
    _section(ws, "Signallar kunlar bo'yicha")
    _header(ws, ["Sana", "Past", "O'rta", "Yuqori", "Jami"])
    for day in a.security.by_day:
        ws.append([day.date, day.past, day.orta, day.yuqori, day.total])
    _section(ws, "Modullar")
    _header(ws, ["Modul", "Signallar", "Ulushi, %", "Tasdiqlangan", "Rad etilgan", "Ko'rilmagan", "Aniqlik, %"])
    for module in a.security.top_modules:
        ws.append([
            f"#{module.code} {module.name}", module.count, module.share, module.confirmed,
            module.rejected, module.unreviewed, _pct(module.precision),
        ])
    _section(ws, "Kameralar")
    _header(ws, ["Kamera", "Bino", "Signallar", "Ulushi, %"])
    for camera in a.security.top_cameras:
        ws.append([camera.name, camera.building, camera.count, camera.share])
    _widths(ws, [40, 22, 14, 14, 14, 14, 14])

    ws = wb.create_sheet("Darslar")
    _section(ws, "Darslar kunlar bo'yicha")
    _header(ws, ["Sana", "Darslar", "O'rtacha diqqat, %", "Uyqu holatlari"])
    for day in a.lessons.by_day:
        ws.append([day.date, day.sessions, _pct(day.attention), day.sleep])
    _widths(ws, [16, 12, 20, 16])

    ws = wb.create_sheet("Tizim")
    _section(ws, "Kameralar (hozirgi holat)")
    _header(ws, ["Jami", "Faol", "Aloqada", "Aloqada, %"])
    ws.append([a.system.cameras_total, a.system.cameras_active, a.system.cameras_live, _pct(a.system.live_rate)])
    _section(ws, "Yuzni tasdiqlash qamrovi")
    _header(ws, ["Guruh", "Jami", "Tasdiqlangan", "Qamrov, %"])
    for row in a.system.coverage:
        ws.append([row.label, row.total, row.confirmed, _pct(row.percent)])
    _widths(ws, [22, 14, 14, 14])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
